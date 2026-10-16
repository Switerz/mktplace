"""DRE-CMP-SKU-1 — custo unitario CMP por SKU a partir das planilhas fiscais (dry-run).

O QUE ELE FAZ
-------------
Le, na pasta indicada, SOMENTE as planilhas do conjunto autorizado e extrai o
CMP unitario por SKU da filial SP:

  * comparativos AP, BB, Kokeshi e MAGA — aba `Consolidado CMP`, `Filial = SP`;
  * CMV-LCV — aba `Resumo Produto`, coluna `CMP Unit`, `Filial = 0301` (SP).

Cada linha candidata carrega competencia/vigencia, entidade, filial, SKU,
descricao auxiliar, CMP e procedencia (arquivo, aba, linha, coluna, sha256).

O QUE ELE NAO FAZ
-----------------
  * nao le receita, canal, `MKP`, CMV realizado nem bonificacao — as abas
    `CMV Jul-26`, `Bonificacoes`, `Planilha1`, `Itens Detalhados`, `Compras`
    nunca sao abertas;
  * nao troca CMP por CMM, SB2, D2, custo de compra ou qualquer outra coluna:
    sem a coluna CMP o arquivo e' RECUSADO;
  * nao resolve conflito por media, maximo, minimo ou "ultimo": a chave com
    mais de um CMP fica marcada como conflito, sem valor;
  * nao escreve em banco nenhum. Nao ha modo de escrita neste gate.

REGIME TEMPORAL OBSERVADO
-------------------------
Nos comparativos o CMP e' uma serie MENSAL (`Jan/25` ... `Jul/26`): o valor e'
o CMP vigente no fim do mes e repete enquanto nao ha nova entrada. Antes da
primeira entrada a celula vem `0` — isso e' AUSENCIA de CMP, nao custo zero.
`Custo Atual` e' a fotografia na data de corte do arquivo. No CMV-LCV ha uma
unica fotografia da competencia do nome do arquivo.

USO
---
    python -m pipelines.dre.cmp_sku --pasta dre
    python -m pipelines.dre.cmp_sku --pasta dre --com-torre --de 2026-01 --ate 2026-09

Codigos de saida: 0 ok; 65 fonte recusada; 69 Torre indisponivel; 2 argumento.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import math
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Optional

logger = logging.getLogger("pipelines.dre.cmp_sku")

VEREDITO_OK = "CMP_SKU_DRY_RUN_READY"
EXIT_OK = 0
EXIT_RECUSA = 65
EXIT_TORRE = 69

MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
         "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}

#: Conjunto autorizado. Arquivo fora destes padroes nunca e' aberto.
RE_COMPARATIVO = re.compile(
    r"^Comparativo_MesAMes_(?P<ent>AP_COSMETICS|BB_VAREJO|KOKESHI|MAGA_COMERCIO)"
    r"_Consolidado(?:_TodasFiliais)?_(?P<de>\d{8})_(?P<ate>\d{8})\.xlsx$")
RE_LCV = re.compile(r"^CMV-LCV-(?P<mes>[a-z]{3})-(?P<ano>\d{4})\.xlsx$")

ABA_COMPARATIVO = "Consolidado CMP"
ABA_LCV = "Resumo Produto"
FILIAL_SP_COMPARATIVO = "SP"
#: Decisao do responsavel (13/10/2026): no CMV-LCV, `0301` e' a filial SP.
FILIAL_SP_LCV = "0301"

CABECALHO_FIXO = ("Filial", "Empresa", "Cód. Próprio", "GTIN", "Descrição")
CABECALHO_FIM = ("Custo Atual", "Entradas")
RE_MES = re.compile(r"^(Jan|Fev|Mar|Abr|Mai|Jun|Jul|Ago|Set|Out|Nov|Dez)/(\d{2})$")

CABECALHO_LCV = ("Filial", "Produto", "Marca", "Linhas", "Qtd", "CMP Unit", "CMM Unit", "CM1 SB2",
                 "CustUnit D2", "Custo CMP", "Custo CMM", "Custo SB2", "Custo D2", "Venda",
                 "Margem CMM %", "Diagnóstico")
#: As unicas colunas do CMV-LCV que saem do parser. `Venda`, `Custo *` e as
#: demais bases de custo sao validadas no cabecalho e descartadas.
COLUNAS_LCV_LIDAS = ("Filial", "Produto", "Marca", "CMP Unit", "Diagnóstico")

#: Codigo proprio que nao identifica produto (varios itens distintos o usam).
RE_SKU_INVALIDO = re.compile(r"^(|SEM GTIN|SEM CODIGO|SEM CÓDIGO|N/A|-)$", re.IGNORECASE)

REGIME_MENSAL = "mensal"
REGIME_ATUAL = "fotografia_atual"


class FonteRecusada(RuntimeError):
    """Arquivo fora do contrato: nada dele e' usado."""


class ArquivoAlterado(FonteRecusada):
    """O arquivo mudou entre o hash inicial e o final da leitura."""


@dataclass(frozen=True)
class Fonte:
    arquivo: str
    path: Path
    entidade: str
    layout: str            # "comparativo" | "lcv"
    aba: str
    filial_sp: str
    corte: date            # ultimo mes coberto pelo arquivo


@dataclass(frozen=True)
class CmpLinha:
    regime: str
    competencia: Optional[date]   # mes do valor; None na fotografia `Custo Atual`
    vigente_em: date              # competencia, ou o mes de corte na fotografia
    entidade: str
    marca: Optional[str]
    filial: str
    sku: str
    gtin: Optional[str]
    descricao: Optional[str]
    cmp: Decimal
    arquivo: str
    aba: str
    linha: int
    coluna: str
    sha256: str


@dataclass
class Leitura:
    fonte: Fonte
    sha256: str
    linhas: list[CmpLinha] = field(default_factory=list)
    q: Counter = field(default_factory=Counter)
    filiais: Counter = field(default_factory=Counter)
    marcas_sem_depara: Counter = field(default_factory=Counter)
    duplicatas_exatas: int = 0
    skus_invalidos: list[dict] = field(default_factory=list)
    links_externos: list[str] = field(default_factory=list)
    formulas_extraidas: int = 0
    meses: list[date] = field(default_factory=list)
    #: (sku, linha) -> serie bruta, para a analise de carry-forward.
    series: list[tuple[str, list[Optional[Decimal]], Optional[Decimal]]] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Utilitarios
# ---------------------------------------------------------------------------

def norm(valor: Any) -> str:
    if valor is None:
        return ""
    return unicodedata.normalize("NFC", str(valor)).strip()


def chave_sku(valor: Any) -> str:
    """Forma de comparacao de codigo: sem espacos nas pontas, caixa alta."""
    return norm(valor).upper()


def sha256_arquivo(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for bloco in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloco)
    return h.hexdigest()


def classificar_valor(v: Any) -> tuple[str, Optional[Decimal]]:
    """`positivo` | `zero` | `negativo` | `nulo` | `nao_numerico`.

    Texto nunca e' convertido ("1,23" fica `nao_numerico`): a fonte entrega
    numero, e um texto ali e' defeito a ser visto, nao adivinhado.
    """
    if v is None or (isinstance(v, str) and not v.strip()):
        return "nulo", None
    if isinstance(v, bool) or not isinstance(v, (int, float, Decimal)):
        return "nao_numerico", None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return "nao_numerico", None
    d = v if isinstance(v, Decimal) else Decimal(str(v))
    if d.is_nan():
        return "nao_numerico", None
    if d > 0:
        return "positivo", d
    if d == 0:
        return "zero", d
    return "negativo", d


def mes_de_rotulo(rotulo: str) -> Optional[date]:
    m = RE_MES.match(rotulo)
    if not m:
        return None
    return date(2000 + int(m.group(2)), MESES[m.group(1).lower()], 1)


def somar_meses(d: date, n: int) -> date:
    t = d.year * 12 + d.month - 1 + n
    return date(t // 12, t % 12 + 1, 1)


# ---------------------------------------------------------------------------
# Conjunto autorizado
# ---------------------------------------------------------------------------

def identificar(path: Path) -> Optional[Fonte]:
    nome = path.name
    m = RE_COMPARATIVO.match(nome)
    if m:
        ate = m.group("ate")
        corte = date(int(ate[4:]), int(ate[2:4]), 1)
        return Fonte(nome, path, m.group("ent"), "comparativo", ABA_COMPARATIVO, FILIAL_SP_COMPARATIVO, corte)
    m = RE_LCV.match(nome)
    if m and m.group("mes") in MESES:
        corte = date(int(m.group("ano")), MESES[m.group("mes")], 1)
        return Fonte(nome, path, "LCV", "lcv", ABA_LCV, FILIAL_SP_LCV, corte)
    return None


def descobrir_fontes(pasta: Path, somente: Optional[list[str]] = None) -> tuple[list[Fonte], list[str]]:
    """Fontes autorizadas presentes na pasta e nomes ignorados (nunca abertos)."""
    if not pasta.is_dir():
        raise FonteRecusada(f"pasta inexistente: {pasta.name}")
    fontes, ignorados = [], []
    for p in sorted(pasta.iterdir()):
        if not p.is_file():
            continue
        f = identificar(p)
        if f is None:
            ignorados.append(p.name)
        else:
            fontes.append(f)
    if somente:
        nomes = {f.arquivo for f in fontes}
        fora = [n for n in somente if n not in nomes]
        if fora:
            raise FonteRecusada(f"arquivo fora do conjunto autorizado ou ausente: {', '.join(fora)}")
        fontes = [f for f in fontes if f.arquivo in somente]
    por_entidade = Counter(f.entidade for f in fontes)
    repetidas = sorted(e for e, n in por_entidade.items() if n > 1)
    if repetidas:
        raise FonteRecusada(f"mais de um arquivo para a mesma entidade: {', '.join(repetidas)}")
    if not fontes:
        raise FonteRecusada("nenhum arquivo do conjunto autorizado na pasta")
    return fontes, ignorados


def links_externos(path: Path) -> list[str]:
    """Nome (sem diretorio) de cada pasta de trabalho externa referenciada."""
    alvos = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if n.startswith("xl/externalLinks/_rels/"):
                xml = z.read(n).decode("utf-8", "replace")
                for alvo in re.findall(r'Target="([^"]+)"', xml):
                    alvos.append(alvo.replace("%20", " ").replace("\\", "/").rsplit("/", 1)[-1])
    return sorted(set(alvos))


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------

def ler(fonte: Fonte) -> Leitura:
    """Le uma fonte com hash antes e depois. Arquivo alterado no meio e' recusado."""
    from openpyxl import load_workbook

    antes = sha256_arquivo(fonte.path)
    leitura = Leitura(fonte=fonte, sha256=antes)
    leitura.links_externos = links_externos(fonte.path)
    valores = load_workbook(fonte.path, read_only=True, data_only=True)
    formulas = load_workbook(fonte.path, read_only=True, data_only=False)
    try:
        if fonte.aba not in valores.sheetnames:
            raise FonteRecusada(f"{fonte.arquivo}: aba {fonte.aba!r} ausente")
        linhas_v = list(valores[fonte.aba].iter_rows(values_only=True))
        linhas_f = list(formulas[fonte.aba].iter_rows(values_only=True))
    finally:
        valores.close()
        formulas.close()
    if fonte.layout == "comparativo":
        _parse_comparativo(leitura, linhas_v, linhas_f)
    else:
        _parse_lcv(leitura, linhas_v, linhas_f)
    depois = sha256_arquivo(fonte.path)
    if depois != antes:
        raise ArquivoAlterado(f"{fonte.arquivo}: arquivo alterado durante a leitura")
    return leitura


def _eh_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _cabecalho(linha: Iterable[Any]) -> list[str]:
    cab = [norm(c) for c in linha]
    while cab and not cab[-1]:
        cab.pop()
    return cab


def validar_cabecalho_comparativo(fonte: Fonte, cab: list[str]) -> list[date]:
    """Devolve os meses da serie. Qualquer desvio do contrato recusa o arquivo."""
    if any("CMM" in c.upper() for c in cab):
        raise FonteRecusada(f"{fonte.arquivo}: coluna CMM no cabecalho; CMM nao substitui CMP")
    if tuple(cab[:len(CABECALHO_FIXO)]) != CABECALHO_FIXO or tuple(cab[-len(CABECALHO_FIM):]) != CABECALHO_FIM:
        raise FonteRecusada(f"{fonte.arquivo}: cabecalho inesperado em {fonte.aba!r}")
    rotulos = cab[len(CABECALHO_FIXO):-len(CABECALHO_FIM)]
    meses = [mes_de_rotulo(r) for r in rotulos]
    if not meses or any(m is None for m in meses):
        raise FonteRecusada(f"{fonte.arquivo}: colunas de mes fora do padrao Mmm/AA")
    for a, b in zip(meses, meses[1:]):
        if b != somar_meses(a, 1):
            raise FonteRecusada(f"{fonte.arquivo}: meses nao consecutivos ({a:%Y-%m} -> {b:%Y-%m})")
    if meses[-1] != fonte.corte:
        raise FonteRecusada(
            f"{fonte.arquivo}: ultimo mes {meses[-1]:%Y-%m} difere do corte do nome {fonte.corte:%Y-%m}")
    return meses


def _parse_comparativo(leitura: Leitura, linhas_v: list[tuple], linhas_f: list[tuple]) -> None:
    fonte = leitura.fonte
    titulo = norm(linhas_v[0][0]) if linhas_v and linhas_v[0] else ""
    if "CMP" not in titulo.upper():
        raise FonteRecusada(f"{fonte.arquivo}: titulo da aba nao declara CMP")
    idx_cab = next((i for i, r in enumerate(linhas_v[:10]) if r and norm(r[0]) == "Filial"), None)
    if idx_cab is None:
        raise FonteRecusada(f"{fonte.arquivo}: cabecalho nao encontrado nas 10 primeiras linhas")
    cab = _cabecalho(linhas_v[idx_cab])
    meses = validar_cabecalho_comparativo(fonte, cab)
    leitura.meses = meses
    i_mes0 = len(CABECALHO_FIXO)
    i_atual = i_mes0 + len(meses)
    vistos: Counter = Counter()
    q = leitura.q

    for i in range(idx_cab + 1, len(linhas_v)):
        r = tuple(linhas_v[i]) + (None,) * max(0, len(cab) - len(linhas_v[i]))
        if all(c is None or norm(c) == "" for c in r[:len(cab)]):
            q["linhas_vazias"] += 1
            continue
        q["linhas_lidas"] += 1
        filial = norm(r[0])
        leitura.filiais[filial] += 1
        if filial != fonte.filial_sp:
            continue
        q["linhas_sp"] += 1
        excel = i + 1
        rf = linhas_f[i] if i < len(linhas_f) else ()
        leitura.formulas_extraidas += sum(1 for c in rf[i_mes0:i_atual + 1] if _eh_formula(c))
        vistos[tuple(norm(c) if isinstance(c, str) else c for c in r[:len(cab)])] += 1
        sku_bruto = norm(r[2])
        if RE_SKU_INVALIDO.match(sku_bruto):
            q["sku_invalido"] += 1
            leitura.skus_invalidos.append({"linha": excel, "sku": sku_bruto, "descricao": norm(r[4])})
            continue
        sku, gtin, desc = chave_sku(sku_bruto), norm(r[3]) or None, norm(r[4]) or None
        base = dict(entidade=fonte.entidade, marca=None, filial=filial, sku=sku, gtin=gtin, descricao=desc,
                    arquivo=fonte.arquivo, aba=fonte.aba, linha=excel, sha256=leitura.sha256)

        serie: list[Optional[Decimal]] = []
        visto_positivo = False
        for j, comp in enumerate(meses):
            estado, valor = classificar_valor(r[i_mes0 + j])
            if estado == "positivo":
                visto_positivo = True
                leitura.linhas.append(CmpLinha(regime=REGIME_MENSAL, competencia=comp, vigente_em=comp, cmp=valor,
                                               coluna=cab[i_mes0 + j], **base))
                serie.append(valor)
                continue
            serie.append(None)
            if estado == "zero":
                # Antes da 1a entrada o zero e' ausencia; depois dela e' anomalia.
                q["zero_antes_primeira_entrada" if not visto_positivo else "zero_apos_cmp"] += 1
            else:
                q[f"cmp_{estado}"] += 1
        estado, atual = classificar_valor(r[i_atual])
        if estado == "positivo":
            leitura.linhas.append(CmpLinha(regime=REGIME_ATUAL, competencia=None, vigente_em=fonte.corte, cmp=atual,
                                           coluna=CABECALHO_FIM[0], **base))
        else:
            q[f"custo_atual_{estado}"] += 1
            atual = None
        leitura.series.append((sku, serie, atual))
    leitura.duplicatas_exatas = sum(n - 1 for n in vistos.values() if n > 1)


def validar_cabecalho_lcv(fonte: Fonte, cab: list[str]) -> None:
    if "CMP Unit" not in cab:
        raise FonteRecusada(f"{fonte.arquivo}: coluna 'CMP Unit' ausente; CMM/SB2/D2 nao substituem CMP")
    if tuple(cab) != CABECALHO_LCV:
        raise FonteRecusada(f"{fonte.arquivo}: cabecalho inesperado em {fonte.aba!r}")


def _parse_lcv(leitura: Leitura, linhas_v: list[tuple], linhas_f: list[tuple]) -> None:
    fonte = leitura.fonte
    if not linhas_v:
        raise FonteRecusada(f"{fonte.arquivo}: aba {fonte.aba!r} vazia")
    cab = _cabecalho(linhas_v[0])
    validar_cabecalho_lcv(fonte, cab)
    ix = {c: cab.index(c) for c in COLUNAS_LCV_LIDAS}
    leitura.meses = [fonte.corte]
    vistos: Counter = Counter()
    q = leitura.q
    for i in range(1, len(linhas_v)):
        r = tuple(linhas_v[i]) + (None,) * max(0, len(cab) - len(linhas_v[i]))
        if all(c is None or norm(c) == "" for c in r[:len(cab)]):
            q["linhas_vazias"] += 1
            continue
        filial = norm(r[ix["Filial"]])
        if filial.upper() == "TOTAL":
            q["linhas_total"] += 1
            continue
        q["linhas_lidas"] += 1
        leitura.filiais[filial] += 1
        if filial != fonte.filial_sp:
            continue
        q["linhas_sp"] += 1
        excel = i + 1
        rf = linhas_f[i] if i < len(linhas_f) else ()
        if ix["CMP Unit"] < len(rf) and _eh_formula(rf[ix["CMP Unit"]]):
            leitura.formulas_extraidas += 1
        lidas = tuple(norm(r[ix[c]]) if isinstance(r[ix[c]], str) else r[ix[c]] for c in COLUNAS_LCV_LIDAS)
        vistos[lidas] += 1
        marca_bruta = norm(r[ix["Marca"]])
        marca = None
        if marca_bruta.isdigit():
            leitura.marcas_sem_depara[marca_bruta] += 1
        elif marca_bruta:
            marca = marca_bruta.lower()
        else:
            q["marca_vazia"] += 1
        sku_bruto = norm(r[ix["Produto"]])
        if RE_SKU_INVALIDO.match(sku_bruto):
            q["sku_invalido"] += 1
            leitura.skus_invalidos.append({"linha": excel, "sku": sku_bruto, "descricao": None})
            continue
        estado, valor = classificar_valor(r[ix["CMP Unit"]])
        q[f"diagnostico_fonte:{norm(r[ix['Diagnóstico']]) or '(vazio)'}"] += 1
        if estado != "positivo":
            # Uma fotografia mensal com 0 nao diz se e' ausencia ou custo zero:
            # nao vira CMP em nenhum dos dois casos.
            q[f"cmp_{estado}"] += 1
            leitura.series.append((chave_sku(sku_bruto), [None], None))
            continue
        leitura.linhas.append(CmpLinha(REGIME_MENSAL, fonte.corte, fonte.corte, fonte.entidade, marca, filial,
                                       chave_sku(sku_bruto), None, None, valor, fonte.arquivo, fonte.aba, excel,
                                       "CMP Unit", leitura.sha256))
        leitura.series.append((chave_sku(sku_bruto), [valor], None))
    leitura.duplicatas_exatas = sum(n - 1 for n in vistos.values() if n > 1)


# ---------------------------------------------------------------------------
# Indice e qualidade
# ---------------------------------------------------------------------------

@dataclass
class Indice:
    """CMP por (entidade, filial, sku). Conjunto de valores => conflito visivel."""
    mensal: dict = field(default_factory=lambda: defaultdict(lambda: defaultdict(set)))
    atual: dict = field(default_factory=lambda: defaultdict(set))
    descricoes: dict = field(default_factory=lambda: defaultdict(set))
    gtins: dict = field(default_factory=lambda: defaultdict(set))
    corte: dict = field(default_factory=dict)          # entidade -> mes de corte
    fotografia: set = field(default_factory=set)       # entidades cuja fonte e' fotografia unica
    por_sku: dict = field(default_factory=lambda: defaultdict(set))  # sku -> {(entidade, filial)}

    def chaves(self) -> set:
        return set(self.mensal) | set(self.atual)


def construir_indice(leituras: list[Leitura]) -> Indice:
    ind = Indice()
    for le in leituras:
        ind.corte[le.fonte.entidade] = le.fonte.corte
        if le.fonte.layout == "lcv":
            ind.fotografia.add(le.fonte.entidade)
        for ln in le.linhas:
            k = (ln.entidade, ln.filial, ln.sku)
            if ln.regime == REGIME_MENSAL:
                ind.mensal[k][ln.competencia].add(ln.cmp)
            else:
                ind.atual[k].add(ln.cmp)
            if ln.descricao:
                ind.descricoes[k].add(ln.descricao)
            if ln.gtin:
                ind.gtins[k].add(ln.gtin)
            ind.por_sku[ln.sku].add((ln.entidade, ln.filial))
    return ind


def unico(valores: set) -> Optional[Decimal]:
    return next(iter(valores)) if len(valores) == 1 else None


def qualidade(leituras: list[Leitura], ind: Indice) -> dict:
    por_fonte = {}
    for le in leituras:
        skus = {ln.sku for ln in le.linhas}
        comps = [ln.competencia for ln in le.linhas if ln.competencia]
        por_fonte[le.fonte.arquivo] = {
            "entidade": le.fonte.entidade,
            "aba": le.fonte.aba,
            "filial_sp": le.fonte.filial_sp,
            "sha256": le.sha256,
            "linhas_lidas": le.q["linhas_lidas"],
            "linhas_sp": le.q["linhas_sp"],
            "linhas_por_filial": dict(sorted(le.filiais.items())),
            "skus_unicos_com_cmp": len(skus),
            "linhas_candidatas": len(le.linhas),
            "duplicatas_exatas": le.duplicatas_exatas,
            "contadores": dict(sorted((k, v) for k, v in le.q.items() if k not in ("linhas_lidas", "linhas_sp"))),
            "marcas_sem_depara": dict(sorted(le.marcas_sem_depara.items())),
            "marca_na_fonte": "codigo/coluna Marca" if le.fonte.layout == "lcv" else "ausente (aba CMP nao tem marca)",
            "competencia_min": min(comps).isoformat() if comps else None,
            "competencia_max": max(comps).isoformat() if comps else None,
            "corte": le.fonte.corte.isoformat(),
            "links_externos_no_arquivo": len(le.links_externos),
            "formulas_nas_celulas_extraidas": le.formulas_extraidas,
            "diferenca_cache_vs_calculo": ("nao_aplicavel: nenhuma formula nas celulas extraidas"
                                           if le.formulas_extraidas == 0 else "nao_verificavel"),
        }
    conflito_mesma_chave = sum(1 for serie in ind.mensal.values() for v in serie.values() if len(v) > 1)
    conflito_atual = sum(1 for v in ind.atual.values() if len(v) > 1)
    sku_desc_multiplas = sum(1 for v in ind.descricoes.values() if len(v) > 1)
    # Mesmo SKU, mesma competencia, entidades diferentes.
    por_sku_comp: dict = defaultdict(lambda: defaultdict(set))
    for (ent, fil, sku), serie in ind.mensal.items():
        for comp, vals in serie.items():
            if len(vals) == 1:
                por_sku_comp[(sku, comp)][ent] |= vals
    multi = [d for d in por_sku_comp.values() if len(d) > 1]
    divergem = [d for d in multi if len({v for s in d.values() for v in s}) > 1]
    skus_multi_ent = {sku for sku, ents in ind.por_sku.items() if len({e for e, _ in ents}) > 1}
    return {
        "por_fonte": por_fonte,
        "consolidado": {
            "linhas_lidas": sum(le.q["linhas_lidas"] for le in leituras),
            "linhas_sp": sum(le.q["linhas_sp"] for le in leituras),
            "linhas_candidatas": sum(len(le.linhas) for le in leituras),
            "chaves_entidade_filial_sku": len(ind.chaves()),
            "skus_unicos": len({k[2] for k in ind.chaves()}),
            "cmp_nulo": sum(le.q["cmp_nulo"] for le in leituras),
            "cmp_zero_antes_primeira_entrada_ausencia": sum(le.q["zero_antes_primeira_entrada"] for le in leituras),
            "cmp_zero_apos_cmp": sum(le.q["zero_apos_cmp"] for le in leituras),
            "cmp_zero_fotografia": sum(le.q["cmp_zero"] for le in leituras),
            "cmp_negativo": sum(le.q["cmp_negativo"] for le in leituras),
            "cmp_nao_numerico": sum(le.q["cmp_nao_numerico"] for le in leituras),
            "duplicatas_exatas": sum(le.duplicatas_exatas for le in leituras),
            "sku_invalido": sum(le.q["sku_invalido"] for le in leituras),
            "conflito_mesmo_sku_mesma_entidade_competencia": conflito_mesma_chave,
            "conflito_custo_atual_mesma_chave": conflito_atual,
            "sku_com_mais_de_uma_descricao_na_entidade": sku_desc_multiplas,
            "skus_em_mais_de_uma_entidade": len(skus_multi_ent),
            "sku_competencia_em_mais_de_uma_entidade": len(multi),
            "sku_competencia_cmp_diferente_entre_entidades": len(divergem),
            "marcas_sem_depara": sum(sum(le.marcas_sem_depara.values()) for le in leituras),
            "links_externos": sum(len(le.links_externos) for le in leituras),
            "formulas_nas_celulas_extraidas": sum(le.formulas_extraidas for le in leituras),
        },
    }


# ---------------------------------------------------------------------------
# Temporalidade
# ---------------------------------------------------------------------------

ALTERNATIVAS = ("por_competencia", "ultimo_disponivel", "entidade_vigencia")


def cmp_alternativa(ind: Indice, chave: tuple, comp: date, alternativa: str) -> tuple[str, Optional[Decimal]]:
    """CMP de uma chave (entidade, filial, sku) para uma venda na competencia `comp`.

    * `por_competencia`: so' o valor da coluna do proprio mes; fora da serie
      ou antes da 1a entrada => ausente.
    * `ultimo_disponivel`: a fotografia `Custo Atual` (ou o ultimo mes, se a
      fonte nao tem fotografia), qualquer que seja o mes da venda.
    * `entidade_vigencia`: intervalos de valor constante por entidade; o
      ultimo intervalo fica aberto adiante do corte ate haver nova entrada.
      Antes da 1a entrada => ausente.

    Devolve (`ok` | `ausente` | `conflito`, valor).
    """
    serie = ind.mensal.get(chave, {})
    if alternativa == "por_competencia":
        vals = serie.get(comp)
    elif alternativa == "ultimo_disponivel":
        vals = ind.atual.get(chave) or (serie[max(serie)] if serie else None)
    elif alternativa == "entidade_vigencia":
        anteriores = [c for c in serie if c <= comp]
        vals = serie[max(anteriores)] if anteriores else None
    else:  # pragma: no cover - programacao
        raise ValueError(alternativa)
    if not vals:
        return "ausente", None
    if len(vals) > 1:
        return "conflito", None
    return "ok", next(iter(vals))


#: Regra decidida em 13/10/2026 (DRE-CMP-SKU-2): alternativa 3.
ALTERNATIVA_DECIDIDA = "entidade_vigencia"


@dataclass(frozen=True)
class Vigencia:
    """CMP vigente de uma chave numa competencia, com a procedencia temporal.

    `competencia_origem` e' o ultimo mes que a FONTE declara ate `competencia_usada`.
    `carry_forward` = o valor esta sendo usado num mes que a fonte nao declara
    (depois do corte da serie, ou depois do mes de uma fotografia).
    """
    estado: str                              # ok | ausente | conflito
    cmp: Optional[Decimal] = None
    competencia_origem: Optional[date] = None
    competencia_usada: Optional[date] = None
    carry_forward: bool = False
    fotografia: bool = False


def cmp_vigente(ind: Indice, chave: tuple, comp: date) -> Vigencia:
    """Alternativa 3 — entidade x vigencia.

    O CMP vale a partir da competencia em que aparece e segue vigente ate nova
    entrada da mesma entidade/filial/SKU. Nunca retroage antes da 1a entrada e
    nunca usa `Custo Atual` (que e' a fotografia do corte, nao uma entrada).
    """
    serie = ind.mensal.get(chave, {})
    anteriores = [c for c in serie if c <= comp]
    if not anteriores:
        return Vigencia("ausente", competencia_usada=comp)
    origem = max(anteriores)
    vals = serie[origem]
    foto = chave[0] in ind.fotografia
    if len(vals) > 1:
        return Vigencia("conflito", None, origem, comp, origem < comp, foto)
    return Vigencia("ok", next(iter(vals)), origem, comp, origem < comp, foto)


def analise_temporal(leituras: list[Leitura], ind: Indice, grade: list[date]) -> dict:
    """Regime observado e quantos SKUs mudam de resultado entre as alternativas."""
    regime = {}
    for le in leituras:
        com_2 = constantes = mudancas = 0
        atual_igual = atual_diff = 0
        for _sku, serie, atual in le.series:
            pos = [v for v in serie if v is not None]
            if len(pos) >= 2:
                com_2 += 1
                trocas = sum(1 for a, b in zip(pos, pos[1:]) if a != b)
                mudancas += trocas
                constantes += trocas == 0
            if atual is not None and serie and serie[-1] is not None:
                if atual == serie[-1]:
                    atual_igual += 1
                else:
                    atual_diff += 1
        regime[le.fonte.arquivo] = {
            "regime": ("serie mensal com valor vigente repetido ate nova entrada (carry-forward)"
                       if le.fonte.layout == "comparativo" else "fotografia unica da competencia do arquivo"),
            "meses_na_serie": len(le.meses),
            "linhas_com_2_ou_mais_meses_com_cmp": com_2,
            "linhas_com_cmp_constante": constantes,
            "trocas_de_valor_na_serie": mudancas,
            "custo_atual_igual_ao_ultimo_mes": atual_igual,
            "custo_atual_diferente_do_ultimo_mes": atual_diff,
        }

    resultado = {a: Counter() for a in ALTERNATIVAS}
    skus_mudam = {"por_competencia_vs_ultimo_disponivel": set(),
                  "por_competencia_vs_entidade_vigencia": set(),
                  "ultimo_disponivel_vs_entidade_vigencia": set()}
    anacronismo = set()
    for chave in sorted(ind.chaves()):
        for comp in grade:
            r = {a: cmp_alternativa(ind, chave, comp, a) for a in ALTERNATIVAS}
            for a, (estado, _) in r.items():
                resultado[a][estado] += 1
            for par in skus_mudam:
                a, b = par.split("_vs_")
                if r[a] != r[b]:
                    skus_mudam[par].add(chave)
            if r["por_competencia"][0] == "ausente" and r["ultimo_disponivel"][0] == "ok":
                anacronismo.add(chave)
    return {
        "regime_por_fonte": regime,
        "grade_meses": [g.isoformat() for g in grade],
        "celulas_chave_mes_por_alternativa": {a: dict(c) for a, c in resultado.items()},
        "chaves_que_mudam_de_resultado": {k: len(v) for k, v in skus_mudam.items()},
        "chaves_com_custo_atribuido_antes_da_1a_entrada_ou_apos_corte_no_ultimo_disponivel": len(anacronismo),
    }


# ---------------------------------------------------------------------------
# Saidas locais
# ---------------------------------------------------------------------------

CAMPOS_CSV = ("regime", "competencia", "vigente_em", "entidade", "marca", "filial", "sku", "gtin", "descricao",
              "cmp", "arquivo", "aba", "linha", "coluna", "sha256")


def escrever_candidatos(path: Path, leituras: list[Leitura]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(CAMPOS_CSV)
        for le in leituras:
            for ln in le.linhas:
                w.writerow([getattr(ln, c) if not isinstance(getattr(ln, c), date) else getattr(ln, c).isoformat()
                            for c in CAMPOS_CSV])


def _json_default(o: Any) -> Any:
    if isinstance(o, (date,)):
        return o.isoformat()
    if isinstance(o, Decimal):
        return str(o)
    if isinstance(o, set):
        return sorted(o)
    raise TypeError(type(o).__name__)


def resumo_publico(qual: dict, temporal: dict, torre: Optional[dict], ignorados: list[str]) -> dict:
    """O que vai para stdout: so' contagens. Nenhum CMP, receita ou margem."""
    out = {
        "veredito": VEREDITO_OK,
        "modo": "dry-run (sem escrita; nao existe modo de escrita neste gate)",
        "arquivos_ignorados_fora_do_conjunto": ignorados,
        "fontes": {k: {c: v[c] for c in ("entidade", "linhas_lidas", "linhas_sp", "skus_unicos_com_cmp",
                                           "duplicatas_exatas", "sha256")}
                   for k, v in qual["por_fonte"].items()},
        "qualidade": qual["consolidado"],
        "temporalidade": {"chaves_que_mudam_de_resultado": temporal["chaves_que_mudam_de_resultado"]},
    }
    if torre is not None:
        out["torre"] = torre.get("publico", {})
    return out


def _mes_arg(s: str) -> date:
    m = re.fullmatch(r"(\d{4})-(\d{2})", s or "")
    if not m or not 1 <= int(m.group(2)) <= 12:
        raise argparse.ArgumentTypeError(f"competencia invalida {s!r}; use YYYY-MM")
    return date(int(m.group(1)), int(m.group(2)), 1)


def grade_meses(de: date, ate: date) -> list[date]:
    out, d = [], de
    while d <= ate:
        out.append(d)
        d = somar_meses(d, 1)
    return out


def main(argv: Optional[list[str]] = None, *, datamart_engine=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m pipelines.dre.cmp_sku",
                                 description="Dry-run do CMP unitario por SKU (filial SP). Nunca escreve em banco.")
    ap.add_argument("--pasta", required=True, type=Path, help="pasta com as planilhas (obrigatoria e explicita)")
    ap.add_argument("--arquivo", action="append", help="restringe a um arquivo do conjunto autorizado (repetivel)")
    ap.add_argument("--saida", type=Path, default=Path(".local/dre_cmp_sku"),
                    help="pasta do relatorio local (nao versionado)")
    ap.add_argument("--de", type=_mes_arg, default=date(2026, 1, 1), help="1a competencia da analise, YYYY-MM")
    ap.add_argument("--ate", type=_mes_arg, default=date(2026, 7, 1), help="ultima competencia da analise, YYYY-MM")
    ap.add_argument("--com-torre", action="store_true",
                    help="cruza com as vendas da Torre (Data Mart, somente leitura)")
    args = ap.parse_args(argv)
    if args.de > args.ate:
        ap.error("--de posterior a --ate")
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s %(message)s")

    try:
        fontes, ignorados = descobrir_fontes(args.pasta, args.arquivo)
        leituras = [ler(f) for f in fontes]
    except FonteRecusada as exc:
        print(json.dumps({"status": "recusado", "motivo": str(exc)}, ensure_ascii=False))
        return EXIT_RECUSA

    ind = construir_indice(leituras)
    qual = qualidade(leituras, ind)
    grade = grade_meses(args.de, args.ate)
    temporal = analise_temporal(leituras, ind, grade)

    torre = None
    if args.com_torre:
        from pipelines.dre import cmp_sku_torre as T
        try:
            torre = T.executar(ind, grade, engine=datamart_engine)
        except T.TorreIndisponivel as exc:
            print(json.dumps({"status": "torre_indisponivel", "detalhe": str(exc)}, ensure_ascii=False))
            return EXIT_TORRE

    args.saida.mkdir(parents=True, exist_ok=True)
    escrever_candidatos(args.saida / "candidatos_cmp.csv", leituras)
    completo = {"qualidade": qual, "temporalidade": temporal, "ignorados": ignorados,
                "links_externos": {le.fonte.arquivo: le.links_externos for le in leituras},
                "skus_invalidos": {le.fonte.arquivo: le.skus_invalidos for le in leituras}}
    if torre is not None:
        # O detalhe por SKU vai so' para o CSV local; o JSON fica agregado.
        completo["torre"] = {k: v for k, v in torre.items() if k != "detalhe"}
        T.escrever_detalhes(args.saida, torre)
    (args.saida / "relatorio.json").write_text(
        json.dumps(completo, ensure_ascii=False, indent=1, default=_json_default), encoding="utf-8")
    print(json.dumps(resumo_publico(qual, temporal, torre, ignorados), ensure_ascii=False, indent=1,
                     default=_json_default))
    return EXIT_OK


if __name__ == "__main__":  # pragma: no cover
    sys.exit(main())
