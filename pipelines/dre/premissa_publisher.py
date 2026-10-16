"""Publisher das premissas percentuais da DRE: matriz XLSX -> Neon (Gate DRE-PREMISSAS-1).

O QUE ELE FAZ
-------------
Le a matriz privada de premissas por marca (uma aba, marcas em C:K, linhas de
DRE na coluna A) e grava em `marts.fact_dre_premissa_mensal` SO' as taxas das
cinco marcas da Torre, para UMA competencia declarada no comentario da matriz.

CONTRATO DA MATRIZ (falha fechada em qualquer desvio)
------------------------------------------------------
  * SHA-256 igual ao `--expected-sha256`, antes e depois da leitura;
  * uma unica aba; comentario em C1 declarando arquivo de origem, data de
    envio e a coluna de referencia (`out/26`), que tem de ser a `--competencia`;
  * marcas conhecidas em C1:K1: as cinco da Torre obrigatorias e uma unica vez;
    as quatro fora do escopo (BySamia, Denavita, Aua, Yenzah) sao ignoradas e
    contadas; qualquer outra marca recusa o arquivo;
  * as 18 linhas de premissa presentes uma unica vez, pelo ROTULO (nunca pela
    cor, nunca pela posicao);
  * celula consumida: numero entre 0 e 1, ou vazia (= ausencia, nunca zero).
    Formula, texto ou booleano recusam o arquivo. A coluna B (formulas do
    modelo) nunca e' lida, e nenhum cache de formula e' usado;
  * linhas que a Torre alimenta (receita, CMV, midia, influencer, marketplace)
    e subtotais tem de estar vazias em C:K: valor ali seria uma premissa que o
    contrato nao conhece.

CONTRATO DE ESCRITA
-------------------
  * dry-run por padrao; escrever exige `--apply`;
  * uma transacao no Neon com `pg_try_advisory_xact_lock` — sem espera, sem retry;
  * backup das linhas do escopo ANTES do DELETE, num arquivo privado fora do Git;
  * DELETE SO' de `competencia x cinco marcas`; contagem fora do escopo
    conferida antes e depois — se mudar, rollback;
  * reconciliacao depois da escrita (o que foi lido de volta = o que foi montado);
  * auditoria em `audit.source_sync_run`; o `sync_run_id` e' o `run_id`;
  * manifesto privado SEM taxas, sem caminho local e sem credencial.

Codigos de saida: 0 ok; 65 matriz recusada; 69 banco indisponivel; 75 lock
ocupado; 78 migration ausente; 1 erro inesperado.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable, Optional

logger = logging.getLogger("pipelines.dre.premissa_publisher")

SOURCE_NAME = "dre_premissa_mensal"
TABELA = "marts.fact_dre_premissa_mensal"
#: Familia 9181400xx (DRE); 026 = numero da migration que cria a tabela.
LOCK_KEY = 918140026

EXIT_OK = 0
EXIT_RECUSA = 65
EXIT_BANCO = 69
EXIT_LOCK = 75
EXIT_MIGRATION = 78

#: Rotulo da coluna C:K -> brand_key. Cinco obrigatorias.
MARCAS_TORRE = {"kokeshi": "kokeshi", "apice": "apice", "barbours": "barbours",
                "rituaria": "rituaria", "lescent": "lescent"}
#: Fora do escopo atual da Torre: ignoradas e contadas, nunca carregadas nem zeradas.
MARCAS_FORA_DO_ESCOPO = frozenset({"bysamia", "denavita", "aua", "yenzah"})

#: Rotulo normalizado da coluna A -> componente. Igual a lista FECHADA da migration 026.
LINHAS_PREMISSA = {
    "devolucoes e reembolsos": "devolucoes_reembolsos",
    "icms": "icms",
    "difal": "difal",
    "pis": "pis",
    "cofins": "cofins",
    "iss/ipi/simples nacional": "iss_ipi_simples",
    "outras receitas operacionais": "outras_receitas_operacionais",
    "frete primario": "frete_primario",
    "bonificacoes": "bonificacoes",
    "embalagens": "embalagens",
    "provisao perda de estoque": "provisao_perda_estoque",
    "publicidade e propaganda": "publicidade_propaganda",
    "comissoes": "comissoes",
    "agencia de growth": "agencia_growth",
    "despesas de gateway": "despesas_gateway",
    "brindes": "brindes",
    "logistica - fretes": "logistica_fretes",
    "logistica - operador": "logistica_operador",
}
COMPONENTES = tuple(LINHAS_PREMISSA.values())
#: Armazenados e visiveis, mas FORA da soma, com o status que explica por que.
#: Igual a `NAO_SOMAVEIS` da migration 026 (CHECK `ck_fdpm_nao_somavel`).
#:   * outras receitas: rotulo de receita, formula de deducao (sinal nao ratificado);
#:   * comissoes: pode se sobrepor as taxas reais de marketplace (nao ratificado).
NAO_SOMAVEIS = {
    "outras_receitas_operacionais": "sinal_nao_ratificado",
    "comissoes": "sobreposicao_nao_ratificada",
}
MESES_NOME = {1: "janeiro", 2: "fevereiro", 3: "marco", 4: "abril", 5: "maio", 6: "junho", 7: "julho",
              8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro"}


def rotulo_publico(competencia: date) -> str:
    """Unica procedencia textual que sai do arquivo: nunca o nome do orcamento."""
    return f"Premissas de {MESES_NOME[competencia.month]}/{competencia.year}"

#: Comentario de C1: `Fonte: <arquivo>, enviado em <AAAA-MM-DD>. ... coluna <mmm>/<aa>.`
RE_FONTE = re.compile(r"Fonte:\s*([^,\s][^,]*?\.xlsx)", re.IGNORECASE)
RE_ENVIO = re.compile(r"enviado em\s*(\d{4}-\d{2}-\d{2})", re.IGNORECASE)
RE_COLUNA = re.compile(r"coluna\s+([a-z]{3})/(\d{2})", re.IGNORECASE)
MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
         "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}


class MatrizRecusada(RuntimeError):
    """A matriz saiu do contrato: nada dela e' usado."""


class ArquivoAlterado(MatrizRecusada):
    pass


class LockOcupado(RuntimeError):
    pass


class MigrationAusente(RuntimeError):
    pass


class EscritaRecusada(RuntimeError):
    """Reconciliacao da escrita falhou: a transacao inteira volta."""


# ---------------------------------------------------------------------------
# Leitura da matriz — PURA (so' o arquivo)
# ---------------------------------------------------------------------------

def sha256_arquivo(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for bloco in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloco)
    return h.hexdigest()


def normalizar_rotulo(v: Any) -> str:
    """`(-) Devoluções e Reembolsos` -> `devolucoes e reembolsos`."""
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    s = re.sub(r"^\(\s*[-=+]\s*\)\s*", "", s.strip())
    return re.sub(r"\s+", " ", s).lower()


def chave_marca(v: Any) -> str:
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


@dataclass
class Matriz:
    competencia: date
    sha256: str
    source_period: str          # ex.: "out/26"
    source_note: str            # rotulo publico: "Premissas de outubro/2026"
    source_sent_on: date        # data de envio declarada no comentario de C1
    #: Nome do arquivo de orcamento declarado em C1. SO' no manifesto privado:
    #: nunca vai para o banco, a API ou o stdout.
    arquivo_origem_privado: str = ""
    linhas: list[dict] = field(default_factory=list)   # competencia, brand, component, rate, somavel, status
    marcas_ignoradas: list[str] = field(default_factory=list)


def _competencia_do_comentario(texto: str) -> tuple[date, str, date, str]:
    col = RE_COLUNA.search(texto)
    fonte = RE_FONTE.search(texto)
    envio = RE_ENVIO.search(texto)
    if not (col and fonte and envio) or col.group(1).lower() not in MESES:
        raise MatrizRecusada("comentario de C1 nao declara fonte, data de envio e coluna mmm/aa")
    comp = date(2000 + int(col.group(2)), MESES[col.group(1).lower()], 1)
    nome = fonte.group(1).replace("\\", "/").rsplit("/", 1)[-1].strip()
    return comp, f"{col.group(1).lower()}/{col.group(2)}", date.fromisoformat(envio.group(1)), nome


def _taxa(v: Any, onde: str) -> Optional[Decimal]:
    """Numero entre 0 e 1, ou None (vazio = ausencia). Qualquer outra coisa recusa."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str) and v.startswith("="):
        raise MatrizRecusada(f"{onde}: formula numa celula consumida")
    if isinstance(v, bool) or not isinstance(v, (int, float, Decimal)):
        raise MatrizRecusada(f"{onde}: valor nao numerico")
    d = Decimal(str(v))
    if d.is_nan() or d.is_infinite() or not (Decimal(0) <= d <= Decimal(1)):
        raise MatrizRecusada(f"{onde}: taxa fora de 0..1")
    return d


def ler_matriz(path: Path, *, competencia: date, expected_sha256: str,
               _apos_leitura: Optional[Callable[[], None]] = None) -> Matriz:
    """Le e valida a matriz. `_apos_leitura` existe so' para teste (simula alteracao)."""
    from openpyxl import load_workbook

    if not path.is_file():
        raise MatrizRecusada("arquivo de premissas nao encontrado")
    antes = sha256_arquivo(path)
    if antes != expected_sha256.lower():
        raise MatrizRecusada("SHA-256 da matriz difere do esperado")
    # data_only=False: a celula traz a FORMULA, se houver — e e' recusada. Nunca o cache.
    wb = load_workbook(path, data_only=False)
    try:
        if len(wb.worksheets) != 1:
            raise MatrizRecusada(f"a matriz deveria ter 1 aba, tem {len(wb.worksheets)}")
        ws = wb.worksheets[0]
        comentario = ws["C1"].comment
        if comentario is None:
            raise MatrizRecusada("C1 sem comentario de fonte e competencia")
        comp, periodo, enviado_em, arquivo_origem = _competencia_do_comentario(comentario.text)
        if comp != competencia:
            raise MatrizRecusada(
                f"a matriz declara a competencia {comp:%Y-%m}; pedida {competencia:%Y-%m}")
        # --- marcas em C1:K1 -------------------------------------------------
        colunas, ignoradas = {}, []
        for col in range(3, 12):
            k = chave_marca(ws.cell(1, col).value)
            if not k:
                continue
            if k in MARCAS_FORA_DO_ESCOPO:
                ignoradas.append(k)
            elif k in MARCAS_TORRE:
                if MARCAS_TORRE[k] in colunas.values():
                    raise MatrizRecusada(f"marca duplicada no cabecalho: {k}")
                colunas[col] = MARCAS_TORRE[k]
            else:
                raise MatrizRecusada(f"marca inesperada no cabecalho: {k}")
        faltam = sorted(set(MARCAS_TORRE.values()) - set(colunas.values()))
        if faltam:
            raise MatrizRecusada(f"marca(s) da Torre ausente(s) na matriz: {', '.join(faltam)}")
        if ws.max_column > 11 and any(ws.cell(1, c).value is not None for c in range(12, ws.max_column + 1)):
            raise MatrizRecusada("coluna de marca alem de K")
        # --- linhas -------------------------------------------------------------
        vistas: dict[str, int] = {}
        linhas: list[dict] = []
        for r in range(2, ws.max_row + 1):
            rotulo = normalizar_rotulo(ws.cell(r, 1).value)
            comp_nome = LINHAS_PREMISSA.get(rotulo)
            if comp_nome is None:
                # Linha que a Torre alimenta, subtotal ou titulo: C:K tem de estar vazio.
                if any(ws.cell(r, c).value is not None for c in range(3, 12)):
                    raise MatrizRecusada(f"linha {r} ({rotulo or 'sem rotulo'}) tem valor em C:K fora do contrato")
                continue
            if comp_nome in vistas:
                raise MatrizRecusada(f"linha de premissa duplicada: {comp_nome} (linhas {vistas[comp_nome]} e {r})")
            vistas[comp_nome] = r
            for col in range(3, 12):
                v = ws.cell(r, col).value
                if isinstance(v, str) and v.startswith("="):
                    raise MatrizRecusada(f"{comp_nome}/{ws.cell(1, col).value}: formula numa celula consumida")
            for col, brand in colunas.items():
                taxa = _taxa(ws.cell(r, col).value, f"{comp_nome}/{brand}")
                status = ("ausente_na_origem" if taxa is None else NAO_SOMAVEIS.get(comp_nome, "ok"))
                linhas.append({"competencia": comp, "brand": brand, "component": comp_nome, "rate": taxa,
                               "somavel": comp_nome not in NAO_SOMAVEIS, "status": status})
        ausentes = [c for c in COMPONENTES if c not in vistas]
        if ausentes:
            raise MatrizRecusada(f"linha(s) de premissa ausente(s): {', '.join(ausentes)}")
    finally:
        wb.close()
    if _apos_leitura is not None:
        _apos_leitura()
    if sha256_arquivo(path) != antes:
        raise ArquivoAlterado("a matriz mudou durante a leitura")
    linhas.sort(key=lambda x: (x["brand"], COMPONENTES.index(x["component"])))
    return Matriz(competencia=comp, sha256=antes, source_period=periodo, source_note=rotulo_publico(comp),
                  source_sent_on=enviado_em, arquivo_origem_privado=arquivo_origem,
                  linhas=linhas, marcas_ignoradas=sorted(ignoradas))


# ---------------------------------------------------------------------------
# Manifesto — sem taxas, sem caminho, sem credencial
# ---------------------------------------------------------------------------

def manifesto(m: Matriz, *, modo: str, extra: Optional[dict] = None, privado: bool = False) -> dict:
    """Sem taxas, caminho ou credencial. `privado=True` acrescenta o nome do
    arquivo de orcamento — so' para o arquivo de manifesto fora do Git."""
    por_status: dict[str, int] = {}
    for ln in m.linhas:
        por_status[ln["status"]] = por_status.get(ln["status"], 0) + 1
    return {
        "gate": "DRE-PREMISSAS-1", "modo": modo,
        "competencia": m.competencia.isoformat()[:7], "source_period": m.source_period,
        "source_sha256": m.sha256, "source_note": m.source_note,
        "source_sent_on": m.source_sent_on.isoformat(),
        **({"arquivo_origem_privado": m.arquivo_origem_privado} if privado else {}),
        "marcas_carregadas": sorted({ln["brand"] for ln in m.linhas}),
        "marcas_ignoradas": m.marcas_ignoradas,
        "componentes": list(COMPONENTES), "linhas": len(m.linhas),
        "linhas_por_status": dict(sorted(por_status.items())),
        "zeros_explicitos": sum(1 for ln in m.linhas if ln["rate"] == 0),
        "ausentes": sum(1 for ln in m.linhas if ln["rate"] is None),
        **(extra or {}),
    }


# ---------------------------------------------------------------------------
# Escrita — Neon
# ---------------------------------------------------------------------------

SQL_TABELA_EXISTE = f"SELECT to_regclass('{TABELA}') IS NOT NULL"
SQL_LOCK = f"SELECT pg_try_advisory_xact_lock({LOCK_KEY})"
SQL_ESCOPO = f"""
    SELECT competencia, brand, component, rate, somavel, status, source_period, source_sha256,
           source_note, source_sent_on, loaded_at, run_id
    FROM {TABELA} WHERE competencia = :competencia AND brand = ANY(:marcas)
    ORDER BY brand, component
"""
SQL_FORA_DO_ESCOPO = f"""
    SELECT count(*) FROM {TABELA} WHERE NOT (competencia = :competencia AND brand = ANY(:marcas))
"""
SQL_APAGA_ESCOPO = f"DELETE FROM {TABELA} WHERE competencia = :competencia AND brand = ANY(:marcas)"
SQL_INSERE = f"""
    INSERT INTO {TABELA} (competencia, brand, component, rate, somavel, status, source_period,
        source_sha256, source_note, source_sent_on, loaded_at, run_id)
    VALUES (:competencia, :brand, :component, :rate, :somavel, :status, :source_period,
        :source_sha256, :source_note, :source_sent_on, :loaded_at, :run_id)
"""
SQL_AUDIT_ABRE = """
    INSERT INTO audit.source_sync_run (source_name, status, started_at, source_min_date, source_max_date)
    VALUES (:fonte, 'running', :agora, :competencia, :competencia)
    RETURNING sync_run_id
"""
SQL_AUDIT_FECHA = """
    UPDATE audit.source_sync_run
       SET status = 'success', finished_at = :agora, rows_extracted = :extraidas, rows_loaded = :gravadas
     WHERE sync_run_id = :run_id
"""
SQL_AUDIT_FALHA = """
    INSERT INTO audit.source_sync_run
        (source_name, status, started_at, finished_at, error_message, source_min_date, source_max_date)
    VALUES (:fonte, 'failed', :agora, :agora, :erro, :competencia, :competencia)
"""

_SENSIVEIS = (
    re.compile(r"postgres(?:ql)?(?:\+\w+)?://[^\s'\"]*", re.IGNORECASE),
    re.compile(r"\b\d{1,3}(?:\.\d{1,3}){3}\b"),
    re.compile(r"password=\S+", re.IGNORECASE),
    re.compile(r"\bhost=\S+", re.IGNORECASE),
    re.compile(r"\buser=\S+", re.IGNORECASE),
    re.compile(r"[A-Za-z]:[\\/][^\s'\"]*"),
)


def sanitizar(exc: BaseException | str) -> str:
    texto = exc if isinstance(exc, str) else f"{type(exc).__name__}: {exc}"
    for p in _SENSIVEIS:
        texto = p.sub("[REDACTED]", texto)
    return " ".join(texto.split())[:500]


def _marcas() -> list[str]:
    return sorted(MARCAS_TORRE.values())


def _rows(conn, sql: str, params: dict) -> list[dict]:
    from sqlalchemy import text
    return [dict(r) for r in conn.execute(text(sql), params).mappings().all()]


def _chave(r: dict) -> tuple:
    rate = r["rate"]
    return (r["brand"], r["component"], None if rate is None else Decimal(str(rate)).normalize(),
            bool(r["somavel"]), r["status"])


def diagnosticar(engine, m: Matriz) -> dict:
    """Dry-run: le o Neon em transacao READ ONLY e compara. ZERO escrita."""
    from sqlalchemy import text

    p = {"competencia": m.competencia, "marcas": _marcas()}
    with engine.connect() as conn:
        conn.execute(text("SET TRANSACTION READ ONLY"))
        existe = bool(conn.execute(text(SQL_TABELA_EXISTE)).scalar())
        atual = _rows(conn, SQL_ESCOPO, p) if existe else []
        fora = conn.execute(text(SQL_FORA_DO_ESCOPO), p).scalar() if existe else None
        conn.rollback()
    a, n = {_chave(r)[:2]: _chave(r) for r in atual}, {_chave(r)[:2]: _chave(r) for r in m.linhas}
    return {"tabela_existe": existe, "linhas_no_escopo_hoje": len(atual), "linhas_fora_do_escopo": fora,
            "diff": {"novas": len(set(n) - set(a)), "removidas": len(set(a) - set(n)),
                     "alteradas": sum(1 for k in set(a) & set(n) if a[k] != n[k]),
                     "iguais": sum(1 for k in set(a) & set(n) if a[k] == n[k])}}


def _backup(pasta: Path, m: Matriz, atual: list[dict], agora: datetime) -> Path:
    pasta.mkdir(parents=True, exist_ok=True)
    caminho = pasta / f"backup_premissas_{m.competencia:%Y-%m}_{agora:%Y%m%dT%H%M%S}Z.json"
    caminho.write_text(json.dumps(atual, default=str, ensure_ascii=False, indent=1), encoding="utf-8")
    return caminho


def publicar(engine, m: Matriz, *, pasta_privada: Path, agora: Optional[datetime] = None,
             antes_do_commit: Optional[Callable[[Any], None]] = None) -> dict:
    """`--apply`: transacao unica, lock sem espera, backup antes do DELETE, reconciliacao.

    `antes_do_commit` existe so' para teste: injeta falha depois da escrita e
    antes do commit, para provar que o rollback desfaz tudo.
    """
    from sqlalchemy import text

    agora = agora or datetime.now(timezone.utc)
    p = {"competencia": m.competencia, "marcas": _marcas()}
    try:
        with engine.begin() as conn:
            if not conn.execute(text(SQL_LOCK)).scalar():
                raise LockOcupado(f"advisory lock {LOCK_KEY} ocupado por outra publicacao")
            if not conn.execute(text(SQL_TABELA_EXISTE)).scalar():
                raise MigrationAusente(f"{TABELA} nao existe: aplique a migration 026 antes do --apply")
            run_id = conn.execute(text(SQL_AUDIT_ABRE), {"fonte": SOURCE_NAME, "agora": agora,
                                                          "competencia": m.competencia}).scalar_one()
            atual = _rows(conn, SQL_ESCOPO, p)
            backup = _backup(pasta_privada, m, atual, agora)
            fora_antes = conn.execute(text(SQL_FORA_DO_ESCOPO), p).scalar_one()
            apagadas = conn.execute(text(SQL_APAGA_ESCOPO), p).rowcount
            if apagadas != len(atual):
                raise EscritaRecusada(f"DELETE apagou {apagadas} linha(s); o escopo tinha {len(atual)}")
            for ln in m.linhas:
                conn.execute(text(SQL_INSERE), {**ln, "source_period": m.source_period,
                                                "source_sha256": m.sha256, "source_note": m.source_note,
                                                "source_sent_on": m.source_sent_on,
                                                "loaded_at": agora, "run_id": run_id})
            fora_depois = conn.execute(text(SQL_FORA_DO_ESCOPO), p).scalar_one()
            if fora_depois != fora_antes:
                raise EscritaRecusada(f"linhas fora do escopo mudaram: {fora_antes} -> {fora_depois}")
            gravadas = {_chave(r) for r in _rows(conn, SQL_ESCOPO, p)}
            if gravadas != {_chave(r) for r in m.linhas}:
                raise EscritaRecusada("o que foi lido de volta difere do que foi montado")
            fechou = conn.execute(text(SQL_AUDIT_FECHA), {"agora": agora, "extraidas": len(m.linhas),
                                                          "gravadas": len(m.linhas), "run_id": run_id}).rowcount
            if fechou != 1:
                raise EscritaRecusada(f"auditoria do run {run_id} nao fechou")
            if antes_do_commit is not None:
                antes_do_commit(conn)
    except BaseException as exc:
        _registrar_falha(engine, m, exc, agora)
        raise
    return {"run_id": run_id, "apagadas": apagadas, "gravadas": len(m.linhas),
            "fora_do_escopo_preservadas": fora_antes, "backup": backup.name}


def _registrar_falha(engine, m: Matriz, exc: BaseException, agora: datetime) -> None:
    """Auditoria da falha em transacao PROPRIA. Nunca mascara a causa nem tenta de novo."""
    from sqlalchemy import text

    if isinstance(exc, LockOcupado):
        return
    try:
        with engine.begin() as conn:
            conn.execute(text(SQL_AUDIT_FALHA), {"fonte": SOURCE_NAME, "agora": agora, "erro": sanitizar(exc),
                                                  "competencia": m.competencia})
    except Exception as aud:  # noqa: BLE001
        logger.error("auditoria da falha nao registrada: %s", sanitizar(aud))


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _mes(s: str) -> date:
    mm = re.fullmatch(r"(\d{4})-(\d{2})", s or "")
    if not mm or not 1 <= int(mm.group(2)) <= 12:
        raise argparse.ArgumentTypeError(f"competencia invalida {s!r}; use YYYY-MM")
    return date(int(mm.group(1)), int(mm.group(2)), 1)


def _sha(s: str) -> str:
    if not re.fullmatch(r"[0-9a-fA-F]{64}", s or ""):
        raise argparse.ArgumentTypeError("--expected-sha256 precisa ter 64 hexadecimais")
    return s.lower()


def _engine_padrao():
    from sqlalchemy import create_engine

    from pipelines.common.config import settings
    from pipelines.dre.insumo_publisher import url_psycopg2

    if not settings.database_url:
        raise RuntimeError("DATABASE_URL ausente")
    return create_engine(url_psycopg2(settings.database_url), pool_pre_ping=True, future=True)


def main(argv: Optional[list[str]] = None, *, engine_factory: Optional[Callable] = None,
         agora: Optional[datetime] = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m pipelines.dre.premissa_publisher",
                                 description="Premissas percentuais da DRE (estimadas). Dry-run por padrao.")
    ap.add_argument("--arquivo", required=True, type=Path, help="matriz XLSX privada (nunca versionada)")
    ap.add_argument("--competencia", required=True, type=_mes, help="YYYY-MM declarada no comentario de C1")
    ap.add_argument("--expected-sha256", required=True, type=_sha, help="SHA-256 conferido do arquivo")
    ap.add_argument("--manifesto-dir", type=Path, default=Path(".local/dre_premissas"),
                    help="pasta privada (fora do Git) do manifesto e do backup")
    ap.add_argument("--sem-banco", action="store_true", help="dry-run so' do arquivo, sem consultar o Neon")
    ap.add_argument("--apply", action="store_true", help="escreve no Neon; sem ele e' dry-run")
    args = ap.parse_args(argv)
    if args.apply and args.sem_banco:
        ap.error("--apply e --sem-banco sao incompativeis")
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s %(message)s")
    agora = agora or datetime.now(timezone.utc)

    try:
        m = ler_matriz(args.arquivo, competencia=args.competencia, expected_sha256=args.expected_sha256)
    except MatrizRecusada as exc:
        print(json.dumps({"status": "recusado", "motivo": sanitizar(exc)}, ensure_ascii=False))
        return EXIT_RECUSA

    extra: dict = {}
    codigo = EXIT_OK
    try:
        if args.sem_banco:
            extra["banco"] = "nao_consultado"
        else:
            engine = (engine_factory or _engine_padrao)()
            if args.apply:
                extra["escrita"] = publicar(engine, m, pasta_privada=args.manifesto_dir, agora=agora)
            else:
                extra["neon"] = diagnosticar(engine, m)
    except LockOcupado as exc:
        extra["erro"], codigo = sanitizar(exc), EXIT_LOCK
    except MigrationAusente as exc:
        extra["erro"], codigo = sanitizar(exc), EXIT_MIGRATION
    except EscritaRecusada as exc:
        extra["erro"], codigo = sanitizar(exc), EXIT_RECUSA
    except Exception as exc:  # noqa: BLE001 — banco indisponivel vira codigo, nunca DSN no log
        from sqlalchemy.exc import SQLAlchemyError
        if not isinstance(exc, (SQLAlchemyError, RuntimeError)):
            raise
        extra["erro"], codigo = sanitizar(exc), EXIT_BANCO

    modo = "apply" if args.apply else "dry-run"
    extra = {**extra, "codigo_saida": codigo}
    args.manifesto_dir.mkdir(parents=True, exist_ok=True)
    # O nome do arquivo de orcamento so' existe no manifesto PRIVADO, nunca no stdout.
    (args.manifesto_dir / f"manifesto_premissas_{m.competencia:%Y-%m}_{agora:%Y%m%dT%H%M%S}Z.json").write_text(
        json.dumps(manifesto(m, modo=modo, extra=extra, privado=True), ensure_ascii=False, indent=1, default=str),
        encoding="utf-8")
    print(json.dumps(manifesto(m, modo=modo, extra=extra), ensure_ascii=False, indent=1, default=str))
    return codigo


if __name__ == "__main__":  # pragma: no cover
    sys.exit(main())
