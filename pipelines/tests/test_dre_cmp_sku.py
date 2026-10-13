"""DRE-CMP-SKU-1 — parser de CMP por SKU, qualidade, temporalidade e shadow.

Sem banco e sem as planilhas reais: cada teste monta um XLSX sintetico com o
mesmo layout das fontes (`Consolidado CMP` e `Resumo Produto`).
"""
from __future__ import annotations

import ast
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from pipelines.dre import cmp_sku as C
from pipelines.dre import cmp_sku_torre as T

JUL = date(2026, 7, 1)
MESES = ["Mai/26", "Jun/26", "Jul/26"]
NOME_AP = "Comparativo_MesAMes_AP_COSMETICS_Consolidado_TodasFiliais_01012025_31072026.xlsx"
NOME_MAGA = "Comparativo_MesAMes_MAGA_COMERCIO_Consolidado_01012025_31072026.xlsx"
NOME_LCV = "CMV-LCV-jul-2026.xlsx"


def comparativo(pasta: Path, linhas: list[tuple], nome: str = NOME_AP, meses=MESES,
                cabecalho=None, titulo="Comparativo Mês a Mês – CMP por Filial", aba="Consolidado CMP",
                outras_abas=("CMV Jul-26", "Bonificações Jul-26")) -> Path:
    """`linhas`: (filial, codigo, gtin, descricao, [cmp por mes], custo_atual)."""
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    ws.append([titulo])
    ws.append(["AP COSMETICS LTDA – ES + RJ + SP"])
    ws.append(["Período: título antigo"])
    ws.append([])
    ws.append([])
    ws.append(cabecalho or ["Filial", "Empresa", "Cód. Próprio", "GTIN", "Descrição", *meses,
                            "Custo Atual", "Entradas"])
    for filial, cod, gtin, desc, valores, atual in linhas:
        ws.append([filial, f"EMP {filial}", cod, gtin, desc, *valores, atual, 1])
    for a in outras_abas:
        # Abas proibidas carregam valores-armadilha: se forem lidas, os testes acusam.
        wb.create_sheet(a).append(["Filial", "Canal", "Receita Líquida", 999999])
    p = pasta / nome
    wb.save(p)
    return p


def lcv(pasta: Path, linhas: list[tuple], cabecalho=None) -> Path:
    """`linhas`: (filial, produto, marca, cmp_unit, diagnostico)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo Produto"
    ws.append(list(cabecalho or C.CABECALHO_LCV))
    for filial, prod, marca, cmp, diag in linhas:
        ws.append([filial, prod, marca, 1, 10, cmp, 99.0, 98.0, 97.0, 1.0, 2.0, 3.0, 4.0, 123456.0, 0.5, diag])
    ws.append(["TOTAL", "", "", 0, 0, "", "", "", "", 0, 0, 0, 0, 0, 0, ""])
    wb.create_sheet("Itens Detalhados").append(["Filial", "Venda"])
    p = pasta / NOME_LCV
    wb.save(p)
    return p


def ler_pasta(pasta: Path) -> tuple[list[C.Leitura], C.Indice]:
    fontes, _ = C.descobrir_fontes(pasta)
    leituras = [C.ler(f) for f in fontes]
    return leituras, C.construir_indice(leituras)


# ---------------------------------------------------------------------------
# Filial
# ---------------------------------------------------------------------------

def test_filtro_sp_nos_comparativos(tmp_path):
    comparativo(tmp_path, [
        ("SP", "A1", "", "PROD A 200ML", [1, 2, 3], 3),
        ("RJ", "A1", "", "PROD A 200ML", [7, 7, 7], 7),
        ("ES", "B1", "", "PROD B", [5, 5, 5], 5),
    ])
    (le,), ind = ler_pasta(tmp_path)
    assert le.q["linhas_lidas"] == 3 and le.q["linhas_sp"] == 1
    assert dict(le.filiais) == {"SP": 1, "RJ": 1, "ES": 1}
    assert {ln.filial for ln in le.linhas} == {"SP"}
    assert ind.chaves() == {("AP_COSMETICS", "SP", "A1")}
    assert ind.mensal[("AP_COSMETICS", "SP", "A1")][JUL] == {Decimal("3")}


def test_maga_usa_sp_e_exclui_mg1_mg2_rj(tmp_path):
    comparativo(tmp_path, [
        ("MG1", "10330/1", "", "FORMULA 60 CAPS", [11, 11, 11], 11),
        ("MG2", "10330/1", "", "FORMULA 60 CAPS", [12, 12, 12], 12),
        ("RJ", "10330/1", "", "FORMULA 60 CAPS", [13, 13, 13], 13),
        ("SP", "10330/1", "", "FORMULA 60 CAPS", [10, 10, 10], 10),
    ], nome=NOME_MAGA)
    (le,), ind = ler_pasta(tmp_path)
    assert le.fonte.entidade == "MAGA_COMERCIO"
    assert {ln.cmp for ln in le.linhas} == {Decimal("10")}
    assert ind.chaves() == {("MAGA_COMERCIO", "SP", "10330/1")}


def test_lcv_usa_somente_0301(tmp_path):
    lcv(tmp_path, [("0301", "LC02002", "", 22.5, "OK"), ("0302", "LC02002", "1003", 30.0, "OK")])
    (le,), ind = ler_pasta(tmp_path)
    assert dict(le.filiais) == {"0301": 1, "0302": 1}
    assert le.q["linhas_total"] == 1
    assert [(ln.filial, ln.cmp, ln.competencia) for ln in le.linhas] == [("0301", Decimal("22.5"), JUL)]
    assert ind.chaves() == {("LCV", "0301", "LC02002")}
    # O codigo de marca da 0302 nem e' contado: a linha saiu no filtro de filial.
    assert le.marcas_sem_depara == {}


def test_lcv_conta_codigo_de_marca_sem_depara(tmp_path):
    lcv(tmp_path, [("0301", "LC1", "1007", 5.0, "OK")])
    (le,), _ = ler_pasta(tmp_path)
    assert le.marcas_sem_depara == {"1007": 1}
    assert le.linhas[0].marca is None


# ---------------------------------------------------------------------------
# CMP e nada alem de CMP
# ---------------------------------------------------------------------------

def test_lcv_sem_coluna_cmp_e_recusado_mesmo_com_cmm(tmp_path):
    cab = [c if c != "CMP Unit" else "CMM Unit 2" for c in C.CABECALHO_LCV]
    lcv(tmp_path, [("0301", "LC1", "", 5.0, "OK")], cabecalho=cab)
    with pytest.raises(C.FonteRecusada, match="CMP Unit"):
        ler_pasta(tmp_path)


def test_comparativo_com_coluna_cmm_e_recusado(tmp_path):
    cab = ["Filial", "Empresa", "Cód. Próprio", "GTIN", "Descrição", *MESES, "Custo Atual CMM", "Entradas"]
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)], cabecalho=cab)
    with pytest.raises(C.FonteRecusada, match="CMM"):
        ler_pasta(tmp_path)


def test_aba_cmm_no_lugar_da_cmp_e_recusada(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)], aba="Consolidado CMM",
                titulo="Comparativo – CMM por Filial")
    with pytest.raises(C.FonteRecusada, match="aba 'Consolidado CMP' ausente"):
        ler_pasta(tmp_path)


@pytest.mark.parametrize("cab, motivo", [
    (["Filial", "Empresa", "GTIN", "Cód. Próprio", "Descrição", *MESES, "Custo Atual", "Entradas"],
     "cabecalho inesperado"),
    (["Filial", "Empresa", "Cód. Próprio", "GTIN", "Descrição", "Mai/26", "Jul/26", "Custo Atual", "Entradas"],
     "nao consecutivos"),
    (["Filial", "Empresa", "Cód. Próprio", "GTIN", "Descrição", "Abr/26", "Mai/26", "Jun/26", "Custo Atual",
      "Entradas"], "difere do corte"),
    (["Filial", "Empresa", "Cód. Próprio", "GTIN", "Descrição", "Maio/26", "Jun/26", "Jul/26", "Custo Atual",
      "Entradas"], "padrao Mmm/AA"),
])
def test_cabecalho_inesperado_recusa(tmp_path, cab, motivo):
    n = len(cab) - 7
    comparativo(tmp_path, [("SP", "A1", "", "X", [1] * n, 1)], cabecalho=cab)
    with pytest.raises(C.FonteRecusada, match=motivo):
        ler_pasta(tmp_path)


def test_abas_proibidas_nunca_entram(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)])
    (le,), _ = ler_pasta(tmp_path)
    assert {ln.aba for ln in le.linhas} == {"Consolidado CMP"}
    assert all(ln.cmp != Decimal("999999") for ln in le.linhas)


def test_nao_ha_escrita_nem_modo_apply():
    """O gate nao tem modo de escrita: nem SQL de escrita, nem flag."""
    for mod in (C, T):
        fonte = Path(mod.__file__).read_text(encoding="utf-8")
        arvore = ast.parse(fonte)
        textos = [n.value.upper() for n in ast.walk(arvore) if isinstance(n, ast.Constant) and isinstance(n.value, str)]
        assert not any(k in t for t in textos for k in ("INSERT ", "UPDATE ", "DELETE ", "TRUNCATE", "CREATE "))
    with pytest.raises(SystemExit):
        C.main(["--pasta", ".", "--apply"])


# ---------------------------------------------------------------------------
# Qualidade
# ---------------------------------------------------------------------------

def test_duplicata_exata_e_contada_e_nao_vira_conflito(tmp_path):
    linha = ("SP", "A1", "789", "X", [1, 2, 2], 2)
    comparativo(tmp_path, [linha, linha])
    leituras, ind = ler_pasta(tmp_path)
    q = C.qualidade(leituras, ind)
    assert q["consolidado"]["duplicatas_exatas"] == 1
    assert q["consolidado"]["conflito_mesmo_sku_mesma_entidade_competencia"] == 0
    assert C.cmp_alternativa(ind, ("AP_COSMETICS", "SP", "A1"), JUL, "por_competencia") == ("ok", Decimal("2"))


def test_conflito_na_mesma_chave_nao_e_resolvido(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 2, 2], 2), ("SP", "A1", "", "X", [1, 2, 3], 3)])
    leituras, ind = ler_pasta(tmp_path)
    q = C.qualidade(leituras, ind)
    assert q["consolidado"]["conflito_mesmo_sku_mesma_entidade_competencia"] == 1
    assert q["consolidado"]["conflito_custo_atual_mesma_chave"] == 1
    for alt in C.ALTERNATIVAS:
        assert C.cmp_alternativa(ind, ("AP_COSMETICS", "SP", "A1"), JUL, alt) == ("conflito", None)
    # Junho tem o mesmo valor nas duas linhas: la' nao ha conflito.
    assert C.cmp_alternativa(ind, ("AP_COSMETICS", "SP", "A1"), date(2026, 6, 1), "por_competencia")[0] == "ok"


def test_cmp_diferente_entre_entidades_e_medido(tmp_path):
    comparativo(tmp_path, [("SP", "LC1", "", "X", [5, 5, 5], 5)])
    lcv(tmp_path, [("0301", "LC1", "", 9.0, "OK")])
    leituras, ind = ler_pasta(tmp_path)
    c = C.qualidade(leituras, ind)["consolidado"]
    assert c["skus_em_mais_de_uma_entidade"] == 1
    assert c["sku_competencia_cmp_diferente_entre_entidades"] == 1


def test_nulo_zero_negativo_e_texto(tmp_path):
    comparativo(tmp_path, [
        ("SP", "A1", "", "X", [0, 0, 4], 4),        # zeros antes da 1a entrada = ausencia
        ("SP", "A2", "", "X", [3, 0, 3], 3),        # zero depois de CMP = anomalia
        ("SP", "A3", "", "X", [None, -1, "1,23"], None),
    ])
    (le,), ind = ler_pasta(tmp_path)
    assert le.q["zero_antes_primeira_entrada"] == 2
    assert le.q["zero_apos_cmp"] == 1
    assert le.q["cmp_nulo"] == 1 and le.q["cmp_negativo"] == 1 and le.q["cmp_nao_numerico"] == 1
    assert le.q["custo_atual_nulo"] == 1
    assert all(ln.cmp > 0 for ln in le.linhas)
    assert ("AP_COSMETICS", "SP", "A3") not in ind.chaves()


def test_ausencia_nao_e_zero_no_parser_e_no_shadow(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X 200ML", [0, 0, 0], 0)])
    lcv(tmp_path, [("0301", "LC1", "", 0, "Sem entrada SD1")])
    leituras, ind = ler_pasta(tmp_path)
    assert not ind.chaves()
    lcv_le = next(le for le in leituras if le.fonte.layout == "lcv")
    assert lcv_le.q["cmp_zero"] == 1
    for sku in ("A1", "LC1"):
        a = T.analisar(ind, T.Catalogo(), [venda(sku, "x")])
        s = a["shadow"][0]
        assert s["skus_por_status"] == {"sem_cmp": 1}
        assert s["custo_cmp_coberto"] == 0 and s["receita_coberta"] == 0
        assert s["pct_receita_sem_cmp"] == Decimal("100.00")


def test_sku_placeholder_nao_identifica_produto(tmp_path):
    comparativo(tmp_path, [("SP", "SEM GTIN", "SEM GTIN", "FITA", [1, 1, 1], 1)])
    (le,), ind = ler_pasta(tmp_path)
    assert le.q["sku_invalido"] == 1 and not ind.chaves()


# ---------------------------------------------------------------------------
# Procedencia, hash e conjunto autorizado
# ---------------------------------------------------------------------------

def test_procedencia_completa(tmp_path):
    p = comparativo(tmp_path, [("SP", "A1", "789", "X", [0, 2, 2], 2)])
    (le,), _ = ler_pasta(tmp_path)
    jul = next(ln for ln in le.linhas if ln.competencia == JUL)
    assert (jul.arquivo, jul.aba, jul.linha, jul.coluna) == (p.name, "Consolidado CMP", 7, "Jul/26")
    assert jul.sha256 == C.sha256_arquivo(p)
    atual = next(ln for ln in le.linhas if ln.regime == C.REGIME_ATUAL)
    assert atual.competencia is None and atual.vigente_em == JUL and atual.coluna == "Custo Atual"


def test_arquivo_alterado_durante_a_leitura_e_recusado(tmp_path, monkeypatch):
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)])
    hashes = iter(["a" * 64, "b" * 64])
    monkeypatch.setattr(C, "sha256_arquivo", lambda _p: next(hashes))
    with pytest.raises(C.ArquivoAlterado):
        ler_pasta(tmp_path)


def test_arquivo_fora_do_conjunto_nunca_e_aberto(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)])
    (tmp_path / "DRE mktplace.xlsx").write_bytes(b"nao e xlsx")
    fontes, ignorados = C.descobrir_fontes(tmp_path)
    assert [f.arquivo for f in fontes] == [NOME_AP]
    assert ignorados == ["DRE mktplace.xlsx"]
    with pytest.raises(C.FonteRecusada, match="fora do conjunto"):
        C.descobrir_fontes(tmp_path, ["DRE mktplace.xlsx"])


def test_duas_fontes_da_mesma_entidade_recusam(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)])
    comparativo(tmp_path, [("SP", "A1", "", "X", [1, 1, 1], 1)],
                nome="Comparativo_MesAMes_AP_COSMETICS_Consolidado_01012025_31072026.xlsx")
    with pytest.raises(C.FonteRecusada, match="mesma entidade"):
        C.descobrir_fontes(tmp_path)


def test_main_dry_run_ponta_a_ponta(tmp_path, capsys):
    pasta = tmp_path / "dre"
    pasta.mkdir()
    comparativo(pasta, [("SP", "A1", "", "X", [1, 1, 1], 1)])
    lcv(pasta, [("0301", "LC1", "", 5.0, "OK")])
    saida = tmp_path / "out"
    assert C.main(["--pasta", str(pasta), "--saida", str(saida), "--de", "2026-05", "--ate", "2026-07"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["veredito"] == "CMP_SKU_DRY_RUN_READY"
    assert out["qualidade"]["linhas_sp"] == 2
    assert (saida / "candidatos_cmp.csv").exists() and (saida / "relatorio.json").exists()


def test_main_recusa_com_codigo_65(tmp_path, capsys):
    assert C.main(["--pasta", str(tmp_path), "--saida", str(tmp_path / "o")]) == C.EXIT_RECUSA
    assert json.loads(capsys.readouterr().out)["status"] == "recusado"


# ---------------------------------------------------------------------------
# Temporalidade
# ---------------------------------------------------------------------------

def test_tres_alternativas_temporais(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [0, 2, 3], 3)])
    _, ind = ler_pasta(tmp_path)
    k = ("AP_COSMETICS", "SP", "A1")
    mai, jun, ago = date(2026, 5, 1), date(2026, 6, 1), date(2026, 8, 1)
    # Antes da 1a entrada: so' o "ultimo disponivel" atribui custo (anacronismo).
    assert C.cmp_alternativa(ind, k, mai, "por_competencia") == ("ausente", None)
    assert C.cmp_alternativa(ind, k, mai, "entidade_vigencia") == ("ausente", None)
    assert C.cmp_alternativa(ind, k, mai, "ultimo_disponivel") == ("ok", Decimal("3"))
    # Dentro da serie: competencia e vigencia coincidem; o ultimo nao.
    assert C.cmp_alternativa(ind, k, jun, "por_competencia") == ("ok", Decimal("2"))
    assert C.cmp_alternativa(ind, k, jun, "entidade_vigencia") == ("ok", Decimal("2"))
    assert C.cmp_alternativa(ind, k, jun, "ultimo_disponivel") == ("ok", Decimal("3"))
    # Depois do corte: so' a vigencia aberta e o ultimo cobrem.
    assert C.cmp_alternativa(ind, k, ago, "por_competencia") == ("ausente", None)
    assert C.cmp_alternativa(ind, k, ago, "entidade_vigencia") == ("ok", Decimal("3"))


def test_analise_temporal_conta_chaves_que_mudam(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [0, 2, 3], 3), ("SP", "A2", "", "X", [4, 4, 4], 4)])
    leituras, ind = ler_pasta(tmp_path)
    t = C.analise_temporal(leituras, ind, C.grade_meses(date(2026, 5, 1), date(2026, 8, 1)))
    assert t["chaves_que_mudam_de_resultado"] == {
        "por_competencia_vs_ultimo_disponivel": 2,   # A1 (mai, jun, ago) e A2 (ago)
        "por_competencia_vs_entidade_vigencia": 2,   # so' agosto, nas duas
        "ultimo_disponivel_vs_entidade_vigencia": 1,  # A1 em mai/jun
    }
    reg = t["regime_por_fonte"][NOME_AP]
    assert reg["custo_atual_igual_ao_ultimo_mes"] == 2 and reg["trocas_de_valor_na_serie"] == 1

# ---------------------------------------------------------------------------
# Vigencia (alternativa 3, decidida em 13/10/2026)
# ---------------------------------------------------------------------------

AGO, JUN, MAI = date(2026, 8, 1), date(2026, 6, 1), date(2026, 5, 1)
NOME_BB = "Comparativo_MesAMes_BB_VAREJO_Consolidado_TodasFiliais_01012025_31072026.xlsx"
NOME_KOK = "Comparativo_MesAMes_KOKESHI_Consolidado_TodasFiliais_01012025_31072026.xlsx"


def test_vigencia_nao_retroage_antes_da_primeira_entrada(tmp_path):
    # `Custo Atual` existe (3), mas nao pode cobrir maio, antes da 1a entrada.
    comparativo(tmp_path, [("SP", "A1", "", "X", [0, 2, 3], 3)])
    _, ind = ler_pasta(tmp_path)
    assert C.cmp_vigente(ind, ("AP_COSMETICS", "SP", "A1"), MAI).estado == "ausente"
    v = C.cmp_vigente(ind, ("AP_COSMETICS", "SP", "A1"), JUN)
    assert (v.cmp, v.competencia_origem, v.carry_forward) == (Decimal("2"), JUN, False)


def test_vigencia_carrega_adiante_depois_do_ultimo_mes(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [0, 2, 3], 3)])
    _, ind = ler_pasta(tmp_path)
    v = C.cmp_vigente(ind, ("AP_COSMETICS", "SP", "A1"), AGO)
    assert (v.estado, v.cmp, v.competencia_origem, v.competencia_usada) == ("ok", Decimal("3"), JUL, AGO)
    assert v.carry_forward and not v.fotografia


def test_fotografia_vale_so_a_partir_da_propria_competencia(tmp_path):
    lcv(tmp_path, [("0301", "LC1", "", 22.0, "OK")])
    _, ind = ler_pasta(tmp_path)
    k = ("LCV", "0301", "LC1")
    assert C.cmp_vigente(ind, k, JUN).estado == "ausente"
    jul, ago = C.cmp_vigente(ind, k, JUL), C.cmp_vigente(ind, k, AGO)
    assert (jul.cmp, jul.carry_forward, jul.fotografia) == (Decimal("22"), False, True)
    assert (ago.cmp, ago.competencia_origem, ago.carry_forward, ago.fotografia) == (Decimal("22"), JUL, True, True)


# ---------------------------------------------------------------------------
# Torre: casamento, kits e shadow
# ---------------------------------------------------------------------------

def venda(codigo, canonico, marca="apice", produto="", qtd=10, receita=1000, plataforma="tiktokshop", comp=JUL):
    return {"competencia": comp, "marca": marca, "plataforma": plataforma, "tipo_operacao": "normal",
            "codigo_sku": codigo, "sku_canonico": canonico, "produto": produto,
            "quantidade": Decimal(qtd), "receita": Decimal(receita), "custo_fiscal": Decimal(1), "linhas": 1}


def resolver(ind, codigos, produto="", marca="apice", cat=None, comp=JUL, gtin=True):
    return T.resolver(ind, codigos, comp, produto, cat, marca, T.indice_gtin(ind) if gtin else None)


def test_ean13_valido():
    assert T.ean13_valido("7908787700522")
    assert not T.ean13_valido("7908787700523")
    assert not T.ean13_valido("0000000000029")


def test_colisao_de_codigo_entre_marcas_e_descartada_por_gtin(tmp_path):
    """`10005` e' um condicionador na BB e um perfume Lescent na venda: o GTIN separa."""
    comparativo(tmp_path, [("SP", "10005", "7899999512099", "COND PFM FIN ROSES", [13, 13, 13], 13)], nome=NOME_BB)
    lcv(tmp_path, [("0301", "LC02002", "", 22.0, "OK")])
    _, ind = ler_pasta(tmp_path)
    cat = T.montar_catalogo([{"marca": "Lescent", "sku": "LC02002", "sku_antigo": "10005", "tipo": "PA",
                              "gtin_ean": "7908787702519"}], [])
    r = resolver(ind, {"codigo_sku": "10005", "sku_canonico": "LC02002"}, "DEO COLONIA NO 2 - 100ML",
                 "lescent", cat)
    assert (r.status, r.cmp, r.entidades, r.descartados) == ("coberto", Decimal("22"), ("LCV",), 1)
    assert r.nivel == T.NIVEL_PROVISORIO


def test_codigo_provisorio_escopado_por_marca_e_entidade(tmp_path):
    lcv(tmp_path, [("0301", "LC1", "", 22.0, "OK")])
    comparativo(tmp_path, [("SP", "RT9", "", "CAPS", [5, 5, 5], 5)], nome=NOME_KOK)
    _, ind = ler_pasta(tmp_path)
    ok = resolver(ind, {"codigo_sku": "LC1"}, marca="lescent")
    assert (ok.status, ok.nivel, ok.cmp) == ("coberto", T.NIVEL_PROVISORIO, Decimal("22"))
    # Mesmo codigo, outra marca: a LCV nao esta no escopo da Barbours.
    assert resolver(ind, {"codigo_sku": "LC1"}, marca="barbours").status == "sem_cmp"
    # Codigo existe, mas numa entidade fora do escopo da marca (KOKESHI nao vende Rituaria).
    assert resolver(ind, {"codigo_sku": "RT9"}, marca="rituaria").status == "sem_cmp"
    assert resolver(ind, {"codigo_sku": "LC1"}, marca="marca_sem_escopo").status == "sem_cmp"


def test_dois_candidatos_sem_evidencia_e_mesmo_cmp_sao_ambiguos(tmp_path):
    comparativo(tmp_path, [("SP", "LC1", "", "X", [22, 22, 22], 22)])
    lcv(tmp_path, [("0301", "LC1", "", 22.0, "OK")])
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "LC1"}, marca="lescent")
    assert (r.status, r.cmp) == ("ambiguo", None)


def test_conflito_barbours_preservado_mesmo_com_gtin(tmp_path):
    """Dois codigos do mesmo produto com CMPs diferentes: nem GTIN nem a serie mais recente escolhem."""
    comparativo(tmp_path, [("SP", "10041", "7908787700522", "DEO BD SPL DELIGHT", [8, 8, 8], 8),
                           ("SP", "BB02027", "", "DEO BD SPL MY SWEET DELIGHT 200ML", [0, 0, 9], 9)], nome=NOME_BB)
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "10041", "sku_canonico": "BB02027"},
                 "7908787700522 - Body Splash My Sweet Delight 200ml", "barbours")
    assert (r.status, r.cmp, r.nivel) == ("conflito", None, "")
    assert {(e, s, c) for e, s, c, _ in r.candidatos} == {("BB_VAREJO", "10041", Decimal("8")),
                                                          ("BB_VAREJO", "BB02027", Decimal("9"))}
    # Antes da serie nova existir (maio), o codigo exato e' unico e pode ser usado.
    antes = resolver(ind, {"codigo_sku": "10041", "sku_canonico": "BB02027"},
                     "7908787700522 - Body Splash My Sweet Delight 200ml", "barbours", comp=MAI)
    assert (antes.status, antes.cmp, antes.nivel) == ("coberto", Decimal("8"), T.NIVEL_GTIN)


def test_volume_divergente_descarta_e_concordante_e_evidencia_fraca(tmp_path):
    comparativo(tmp_path, [("SP", "20322", "", "SHAMPOO CCH 1000 ML", [5, 5, 5], 5),
                           ("SP", "20323", "", "SHAMPOO CCH 300 ML", [4, 4, 4], 4)])
    _, ind = ler_pasta(tmp_path)
    assert resolver(ind, {"codigo_sku": "20322"}, "Shampoo Cachos 300 ml", gtin=False).status == "evidencia_diverge"
    ok = resolver(ind, {"codigo_sku": "20323"}, "Shampoo Cachos 300ml")
    assert (ok.status, ok.nivel) == ("coberto", T.NIVEL_VOLUME)


def test_via_gtin_unico_com_volume_compativel(tmp_path):
    comparativo(tmp_path, [("SP", "7898652877919", "7898652877919", "SHAMPOO CCH NUTRITIVO 300ML", [6, 6, 6], 6)])
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "20688", "sku_canonico": "AP01102"},
                 "7898652877919 - (N) Shampoo Cachos Nutritivo 300ml - H1")
    assert (r.status, r.nivel, r.cmp, r.entidades) == ("coberto", T.NIVEL_VIA_GTIN, Decimal("6"), ("AP_COSMETICS",))


def test_via_gtin_exige_volume_que_afaste_reuso(tmp_path):
    comparativo(tmp_path, [("SP", "X1", "7898652877919", "SHAMPOO CCH NUTRITIVO", [6, 6, 6], 6)])
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "20688"}, "7898652877919 - Shampoo Cachos 300ml")
    assert r.status == "sem_cmp" and "gtin_sem_volume_para_afastar_reuso" in r.motivo
    r = resolver(ind, {"codigo_sku": "20688"}, "7898652877919 - Shampoo Cachos 1L")
    assert r.status == "sem_cmp"


def test_via_gtin_reutilizado_com_volume_divergente_fica_sem_casamento(tmp_path):
    """Mesmo EAN em 60 e 180 caps (medido na Rituaria): nenhuma das duas linhas vale."""
    comparativo(tmp_path, [("SP", "10330/1", "7908407007796", "FORMULA 4 MAG 60 CAPS", [11, 11, 11], 11),
                           ("SP", "10369/1", "7908407007796", "FORMULA 4 MAG 180 CAPS", [27, 27, 27], 27)],
                nome="Comparativo_MesAMes_MAGA_COMERCIO_Consolidado_01012025_31072026.xlsx")
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "MKP4MAG"}, "7908407007796 Formula 4mag 60 caps", "rituaria")
    assert r.status == "sem_cmp" and "gtin_reutilizado_volume_incompativel" in r.motivo


def test_via_gtin_ambiguo_fica_sem_casamento(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "7898652877919", "SHAMPOO 300ML", [6, 6, 6], 6),
                           ("SP", "A2", "7898652877919", "SHAMPOO 300 ML", [6, 6, 6], 6)])
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "20688"}, "7898652877919 - Shampoo 300ml")
    assert r.status == "sem_cmp" and "gtin_ambiguo" in r.motivo
    # Venda com dois GTINs validos diferentes tambem e' ambigua.
    r = resolver(ind, {"codigo_sku": "20688"}, "7898652877919 / 7908787700522 Shampoo 300ml")
    assert r.status == "sem_cmp" and "gtin_venda_ambiguo" in r.motivo


def test_codigo_exato_prevalece_sobre_gtin(tmp_path):
    comparativo(tmp_path, [("SP", "20688", "", "SHAMPOO", [5, 5, 5], 5),
                           ("SP", "B9", "7898652877919", "SHAMPOO 300ML", [7, 7, 7], 7)])
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "20688"}, "7898652877919 - Shampoo 300ml")
    assert (r.status, r.cmp, r.nivel) == ("coberto", Decimal("5"), T.NIVEL_PROVISORIO)


def test_via_gtin_nao_resolve_conflito(tmp_path):
    comparativo(tmp_path, [("SP", "C1", "", "X", [5, 5, 5], 5), ("SP", "C2", "", "X", [6, 6, 6], 6),
                           ("SP", "B9", "7898652877919", "SHAMPOO 300ML", [7, 7, 7], 7)])
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": "C1", "sku_canonico": "C2"}, "7898652877919 - Shampoo 300ml")
    assert (r.status, r.cmp) == ("conflito", None)


@pytest.mark.parametrize("codigo, canonico", [("KS03041", "KS03041"), ("KS03042", "KS03011")])
def test_ks03041_e_ks03042_seguem_sem_cmp(tmp_path, codigo, canonico):
    """Mesmo com codigo e GTIN casando na KOKESHI SP, a origem precisa ser confirmada."""
    comparativo(tmp_path, [("SP", codigo, "7908790700687", "SERUM 10ML", [4, 4, 4], 4),
                           ("SP", canonico, "", "X", [4, 4, 4], 4)], nome=NOME_KOK)
    _, ind = ler_pasta(tmp_path)
    r = resolver(ind, {"codigo_sku": codigo, "sku_canonico": canonico}, "7908790700687-SERUM 10ML", "kokeshi")
    assert (r.status, r.cmp) == ("origem_a_confirmar", None)
    a = T.analisar(ind, T.Catalogo(), [venda(codigo, canonico, marca="kokeshi", produto="7908790700687-SERUM 10ML")])
    s = a["shadow"][0]
    assert s["custo_cmp_coberto"] == 0 and s["pct_receita_origem_a_confirmar"] == Decimal("100.00")


def test_kit_sem_regra_nunca_entra_no_shadow(tmp_path):
    comparativo(tmp_path, [("SP", "K1", "", "KIT", [50, 50, 50], 50), ("SP", "C1", "", "X", [10, 10, 10], 10)])
    _, ind = ler_pasta(tmp_path)
    cat = T.montar_catalogo([{"marca": "Ápice", "sku": "K1", "sku_antigo": None, "tipo": "KT", "gtin_ean": None}],
                            [{"kit_codigo": "K1", "componente_codigo": "C1", "qty_per_kit": 2, "kit_marca": "apice"}])
    a = T.analisar(ind, cat, [venda("K1", "K1")])
    s = a["shadow"][0]
    assert s["skus_por_status"] == {"kit_sem_regra": 1}
    assert s["receita_coberta"] == 0 and s["custo_cmp_coberto"] == 0
    assert s["pct_receita_kit_sem_regra"] == Decimal("100.00")
    # As duas leituras existem e foram medidas — e nenhuma foi escolhida.
    assert a["kits"]["ambas_leituras_disponiveis"] == 1


def test_codigo_de_kit_de_outra_marca_nao_torna_o_produto_kit(tmp_path):
    """Medido: `20008` e' kit Barbours no cadastro e hidratante Kokeshi na venda."""
    comparativo(tmp_path, [("SP", "20008", "", "CREME 30G", [4, 4, 4], 4)], nome=NOME_KOK)
    _, ind = ler_pasta(tmp_path)
    cat = T.montar_catalogo([{"marca": "Barbours", "sku": "KBB99045", "sku_antigo": "20008", "tipo": "KT",
                              "gtin_ean": None}], [])
    a = T.analisar(ind, cat, [venda("20008", "KS03015", marca="kokeshi", produto="Creme 30g")])
    assert a["shadow"][0]["skus_por_status"] == {"coberto": 1}


def test_shadow_margem_bruta_apos_cmp_e_carry_forward(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "7898652874765", "COND 500g", [2, 2, 3], 3)])
    _, ind = ler_pasta(tmp_path)
    a = T.analisar(ind, T.Catalogo(), [
        venda("A1", "", produto="7898652874765 - Creme 500g", qtd=10, receita=100),
        venda("SEMCMP", "", qtd=5, receita=50),
        venda("A1", "", produto="7898652874765 - Creme 500g", qtd=1, receita=10, comp=AGO),
    ])
    jul, ago = a["shadow"]
    assert jul["custo_cmp_coberto"] == Decimal("30")            # 10 x CMP de julho (3), nao o de junho
    assert jul["margem_bruta_apos_cmp"] == Decimal("70")
    assert jul["receita_confirmada_gtin"] == Decimal("100")
    assert jul["pct_receita_coberta"] == Decimal("66.67") and jul["pct_receita_sem_cmp"] == Decimal("33.33")
    assert jul["pct_receita_carry_forward"] == Decimal("0.00")
    assert ago["custo_cmp_coberto"] == Decimal("3") and ago["pct_receita_carry_forward"] == Decimal("100.00")
    det = [d for d in a["detalhe"] if d["competencia"] == "2026-08-01"][0]
    assert (det["competencia_origem"], det["competencia_usada"], det["carry_forward"]) == (
        "2026-07-01", "2026-08-01", True)
    assert not any("liquida" in k or "contribuicao" in k for k in jul)


def test_resumo_publico_nao_carrega_valores_monetarios(tmp_path):
    comparativo(tmp_path, [("SP", "A1", "", "X", [2, 2, 3], 3)])
    _, ind = ler_pasta(tmp_path)
    pub = T.resumo_publico(T.analisar(ind, T.Catalogo(), [venda("A1", "")]))
    texto = json.dumps(pub, default=str)
    for proibido in ("receita_fiscal_total", "custo_cmp", "margem", '"cmp"', "receita_confirmada"):
        assert proibido not in texto
