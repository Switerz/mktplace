"""DRE-PREMISSAS-1 — parser da matriz de premissas e CLI, sem banco.

A matriz e' SINTETICA: mesma estrutura da privada (aba unica, marcas em C:K,
comentario em C1, rotulos na coluna A, formulas do modelo na coluna B), com
percentuais inventados. Nenhum percentual real entra no repositorio.
"""
from __future__ import annotations

import ast
import json
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment

from pipelines.dre import premissa_publisher as P

REPO = Path(__file__).resolve().parents[2]
OUT = date(2026, 10, 1)
MARCAS = ["Kokeshi", "Apice", "Barbours", "Rituaria", "Lescent", "BySamia", "Denavita", "Aua", "Yenzah"]
COMENTARIO = ("Fonte: Orcamento_Sintetico_2026.xlsx, enviado em 2026-10-13. % calculadas sobre a Receita "
              "Bruta da DRE - Marketplace de cada marca, coluna out/26.")
#: (rotulo, premissa?) na ordem da matriz real.
ROTULOS = [
    ("Receita Bruta", False), ("(-) Deduções da Receita Bruta", False), ("(-) Devoluções e Reembolsos", True),
    ("(-) Impostos Sob Receita", False), ("ICMS", True), ("DIFAL", True), ("PIS", True), ("COFINS", True),
    ("ISS/IPI/Simples Nacional", True), ("Outras Receitas Operacionais", True), ("Receita Líquida", False),
    ("(-) CMV", False), ("Custo das Mercadorias", False), ("Mark-up direto", False), ("Frete Primário", True),
    ("Bonificações", True), ("Embalagens", True), ("Provisão Perda de Estoque", True), ("Lucro Bruto", False),
    ("Margem Bruta %", False), ("Despesas Comerciais", False), ("Mídia Online", False), ("Influencer", False),
    ("Publicidade e Propaganda", True), ("Comissões", True), ("Agência de Growth", True),
    ("Despesas de Gateway", True), ("Despesas de Marketplace", False), ("Brindes", True),
    ("Despesas Logísticas", False), ("Logística - Fretes", True), ("Logística - Operador", True),
    ("Margem de Contrbuição", False), ("Margem de Contribuição %", False),
]


def taxa_sintetica(i_linha: int, i_marca: int) -> float:
    return round(0.001 * (i_linha + 1) + 0.0001 * i_marca, 6)


def matriz(pasta: Path, *, marcas=None, comentario=COMENTARIO, celulas=None, rotulos=None, abas=1,
           nome="premissas.xlsx") -> tuple[Path, str]:
    """`celulas`: {(rotulo, marca): valor} sobrescreve a taxa sintetica (None = vazio)."""
    marcas = marcas or MARCAS
    wb = Workbook()
    ws = wb.active
    for j, nome_marca in enumerate(marcas):
        ws.cell(1, 3 + j, nome_marca)
    if comentario:
        ws["C1"].comment = Comment(comentario, "autor")
    ws.cell(2, 2, 100)
    for i, (rot, premissa) in enumerate(rotulos or ROTULOS):
        r = i + 2
        ws.cell(r, 1, rot)
        if premissa:
            ws.cell(r, 2, f"=-$B$2*C{r}")
            for j in range(len(marcas)):
                ws.cell(r, 3 + j, taxa_sintetica(i, j))
    for (rot, m), v in (celulas or {}).items():
        r = next(i + 2 for i, (x, _) in enumerate(rotulos or ROTULOS) if x == rot)
        # `ws.cell(r, c, None)` nao limpa a celula no openpyxl: atribuir explicitamente.
        ws.cell(r, 3 + marcas.index(m)).value = v
    for k in range(1, abas):
        wb.create_sheet(f"extra{k}")
    p = pasta / nome
    wb.save(p)
    return p, P.sha256_arquivo(p)


def ler(p, sha, comp=OUT):
    return P.ler_matriz(p, competencia=comp, expected_sha256=sha)


# ---------------------------------------------------------------------------
# Contrato da matriz
# ---------------------------------------------------------------------------

def test_matriz_valida_carrega_cinco_marcas_e_dezoito_componentes(tmp_path):
    m = ler(*matriz(tmp_path))
    assert m.competencia == OUT and m.source_period == "out/26"
    assert m.source_note == "Orcamento_Sintetico_2026.xlsx (enviado em 2026-10-13)"
    assert len(m.linhas) == 5 * 18
    assert {ln["brand"] for ln in m.linhas} == {"kokeshi", "apice", "barbours", "rituaria", "lescent"}
    assert m.marcas_ignoradas == ["aua", "bysamia", "denavita", "yenzah"]
    assert tuple(dict.fromkeys(ln["component"] for ln in m.linhas)) == P.COMPONENTES


def test_componentes_iguais_aos_da_migration():
    src = (REPO / "apps/api/alembic/versions/026_create_fact_dre_premissa_mensal.py").read_text(encoding="utf-8")
    mod = ast.parse(src)
    consts = {t.id: ast.literal_eval(n.value) for n in mod.body if isinstance(n, ast.Assign)
              for t in n.targets if isinstance(t, ast.Name) and t.id in ("COMPONENTES", "MARCAS")}
    assert consts["COMPONENTES"] == P.COMPONENTES
    assert set(consts["MARCAS"]) == set(P.MARCAS_TORRE.values())
    for c in P.COMPONENTES:
        assert f"'{c}'" in src  # o CHECK literal tem a mesma lista


def test_hash_diferente_do_esperado_recusa(tmp_path):
    p, _ = matriz(tmp_path)
    with pytest.raises(P.MatrizRecusada, match="SHA-256"):
        ler(p, "0" * 64)


def test_arquivo_alterado_durante_a_leitura_recusa(tmp_path):
    p, sha = matriz(tmp_path)
    with pytest.raises(P.ArquivoAlterado):
        P.ler_matriz(p, competencia=OUT, expected_sha256=sha,
                     _apos_leitura=lambda: p.write_bytes(p.read_bytes() + b"x"))


def test_celula_vazia_e_ausencia_e_zero_e_premissa_explicita(tmp_path):
    m = ler(*matriz(tmp_path, celulas={("ICMS", "Kokeshi"): None, ("DIFAL", "Kokeshi"): 0}))
    icms = next(ln for ln in m.linhas if ln["brand"] == "kokeshi" and ln["component"] == "icms")
    difal = next(ln for ln in m.linhas if ln["brand"] == "kokeshi" and ln["component"] == "difal")
    assert (icms["rate"], icms["status"]) == (None, "ausente_na_origem")
    assert (difal["rate"], difal["status"]) == (Decimal(0), "ok")


def test_formula_numa_celula_consumida_recusa(tmp_path):
    p, sha = matriz(tmp_path, celulas={("PIS", "Apice"): "=0.01+0.02"})
    with pytest.raises(P.MatrizRecusada, match="formula"):
        ler(p, sha)


def test_formula_em_marca_fora_do_escopo_tambem_recusa(tmp_path):
    p, sha = matriz(tmp_path, celulas={("PIS", "Yenzah"): "=C8"})
    with pytest.raises(P.MatrizRecusada, match="formula"):
        ler(p, sha)


def test_marca_fora_do_escopo_e_ignorada_e_contada(tmp_path):
    m = ler(*matriz(tmp_path))
    assert "bysamia" not in {ln["brand"] for ln in m.linhas}
    assert len(m.marcas_ignoradas) == 4


def test_marca_inesperada_recusa(tmp_path):
    p, sha = matriz(tmp_path, marcas=MARCAS[:8] + ["MarcaNova"])
    with pytest.raises(P.MatrizRecusada, match="marca inesperada"):
        ler(p, sha)


def test_falta_de_uma_das_cinco_marcas_recusa(tmp_path):
    p, sha = matriz(tmp_path, marcas=[m for m in MARCAS if m != "Lescent"])
    with pytest.raises(P.MatrizRecusada, match="lescent"):
        ler(p, sha)


def test_marca_duplicada_recusa(tmp_path):
    p, sha = matriz(tmp_path, marcas=["Kokeshi", "Apice", "Barbours", "Rituaria", "Lescent", "Kokeshi"])
    with pytest.raises(P.MatrizRecusada, match="duplicada"):
        ler(p, sha)


def test_linha_obrigatoria_ausente_recusa(tmp_path):
    p, sha = matriz(tmp_path, rotulos=[r for r in ROTULOS if r[0] != "Brindes"])
    with pytest.raises(P.MatrizRecusada, match="brindes"):
        ler(p, sha)


def test_linha_duplicada_recusa(tmp_path):
    p, sha = matriz(tmp_path, rotulos=ROTULOS + [("COFINS", True)])
    with pytest.raises(P.MatrizRecusada, match="duplicada"):
        ler(p, sha)


def test_valor_em_linha_alimentada_pela_torre_recusa(tmp_path):
    p, sha = matriz(tmp_path, celulas={("Mídia Online", "Kokeshi"): 0.05})
    with pytest.raises(P.MatrizRecusada, match="fora do contrato"):
        ler(p, sha)


@pytest.mark.parametrize("valor", [-0.01, 1.01, "5%", True])
def test_percentual_negativo_acima_de_100_ou_nao_numerico_recusa(tmp_path, valor):
    p, sha = matriz(tmp_path, celulas={("COFINS", "Barbours"): valor})
    with pytest.raises(P.MatrizRecusada):
        ler(p, sha)


def test_competencia_diferente_da_declarada_recusa(tmp_path):
    p, sha = matriz(tmp_path)
    with pytest.raises(P.MatrizRecusada, match="2026-10"):
        ler(p, sha, comp=date(2026, 11, 1))


def test_sem_comentario_ou_com_mais_de_uma_aba_recusa(tmp_path):
    p, sha = matriz(tmp_path, comentario=None)
    with pytest.raises(P.MatrizRecusada, match="comentario"):
        ler(p, sha)
    p, sha = matriz(tmp_path, abas=2, nome="duas.xlsx")
    with pytest.raises(P.MatrizRecusada, match="1 aba"):
        ler(p, sha)


def test_outras_receitas_nunca_somavel(tmp_path):
    m = ler(*matriz(tmp_path))
    outras = [ln for ln in m.linhas if ln["component"] == "outras_receitas_operacionais"]
    assert len(outras) == 5 and all(not ln["somavel"] and ln["status"] == "sinal_nao_ratificado" for ln in outras)
    assert all(ln["somavel"] for ln in m.linhas if ln["component"] != "outras_receitas_operacionais")


def test_coluna_b_nunca_e_lida(tmp_path):
    """Trocar as formulas/valores da coluna B nao muda nada do que e' extraido."""
    p1, s1 = matriz(tmp_path, nome="a.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(p1)
    for r in range(2, 40):
        wb.active.cell(r, 2, 999)
    p2 = tmp_path / "b.xlsx"
    wb.save(p2)
    a, b = ler(p1, s1).linhas, ler(p2, P.sha256_arquivo(p2)).linhas
    assert a == b


# ---------------------------------------------------------------------------
# Manifesto e CLI
# ---------------------------------------------------------------------------

def test_manifesto_sem_taxas_sem_caminho_sem_credencial(tmp_path):
    p, sha = matriz(tmp_path)
    m = ler(p, sha)
    texto = json.dumps(P.manifesto(m, modo="dry-run"), default=str)
    for ln in m.linhas:
        assert str(ln["rate"]) not in texto or ln["rate"] in (None, Decimal(0))
    assert str(tmp_path) not in texto and "\\" not in texto and "postgres" not in texto
    assert not re.search(r"0\.0\d{2,}", texto)


def test_cli_dry_run_sem_banco_nao_escreve_e_grava_manifesto(tmp_path, capsys):
    p, sha = matriz(tmp_path)
    saida = tmp_path / "privado"
    rc = P.main(["--arquivo", str(p), "--competencia", "2026-10", "--expected-sha256", sha,
                 "--manifesto-dir", str(saida), "--sem-banco"],
                agora=datetime(2026, 10, 16, tzinfo=timezone.utc))
    assert rc == 0
    out = json.loads(capsys.readouterr().out)
    assert out["modo"] == "dry-run" and out["linhas"] == 90 and out["banco"] == "nao_consultado"
    assert len(list(saida.glob("manifesto_*.json"))) == 1 and not list(saida.glob("backup_*"))


def test_cli_dry_run_nunca_chama_publicar(tmp_path, monkeypatch, capsys):
    p, sha = matriz(tmp_path)
    monkeypatch.setattr(P, "publicar", lambda *a, **k: (_ for _ in ()).throw(AssertionError("escreveu")))
    monkeypatch.setattr(P, "diagnosticar", lambda engine, m: {"tabela_existe": False})
    rc = P.main(["--arquivo", str(p), "--competencia", "2026-10", "--expected-sha256", sha,
                 "--manifesto-dir", str(tmp_path / "o")], engine_factory=lambda: object())
    assert rc == 0 and json.loads(capsys.readouterr().out)["neon"] == {"tabela_existe": False}


def test_cli_exige_arquivo_competencia_e_sha(tmp_path):
    for faltando in (["--competencia", "2026-10", "--expected-sha256", "a" * 64],
                     ["--arquivo", "x.xlsx", "--expected-sha256", "a" * 64],
                     ["--arquivo", "x.xlsx", "--competencia", "2026-10"]):
        with pytest.raises(SystemExit):
            P.main(faltando)


def test_cli_matriz_recusada_devolve_65(tmp_path, capsys):
    p, _ = matriz(tmp_path)
    rc = P.main(["--arquivo", str(p), "--competencia", "2026-10", "--expected-sha256", "f" * 64,
                 "--manifesto-dir", str(tmp_path / "o"), "--sem-banco"])
    assert rc == P.EXIT_RECUSA and json.loads(capsys.readouterr().out)["status"] == "recusado"


def test_delete_do_publisher_e_sempre_escopado():
    assert re.search(r"DELETE FROM \S+ WHERE competencia = :competencia AND brand = ANY\(:marcas\)",
                     P.SQL_APAGA_ESCOPO)


def test_nenhum_xlsx_ou_percentual_real_versionado():
    """Nada da matriz privada no repositorio: nem o arquivo, nem o nome do orcamento real."""
    import subprocess
    arquivos = subprocess.run(["git", "ls-files"], cwd=REPO, capture_output=True, text=True).stdout.split()
    assert not [a for a in arquivos if a.lower().endswith(".xlsx") and "premiss" in a.lower()]
    assert not [a for a in arquivos if a.startswith("dre/")]
